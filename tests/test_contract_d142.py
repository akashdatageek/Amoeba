"""D142 — the plan–execution contract: typed step checks from a fixed vocabulary, evaluated by plain code in Box 3,
three HARD gate checks in Box 2, and a final requirement check that needs every covering check passed. Offline."""
import json

import pytest

from amoeba.interp.runtime import Interpreter
from amoeba.task import contract as C
from amoeba.task.deliverables import final_check, requirement_check
from amoeba.task.models import DraftPlanStep
from amoeba.task.parsers import parse_plan_d24
from amoeba.task.quality import gate_suggestions
from amoeba.task.schemas import CheckOut, PlanOut, plan_markdown
from tests.test_fixes_d117 import good
from tests.test_plan_runner import plan_team, step_no


def chk(check, covers=("R1",), **params):
    return {"name": check, "params": params, "covers": list(covers)}


# ---- the vocabulary ---------------------------------------------------------------------------------------------
@pytest.mark.parametrize("c,problem", [
    (chk("file_exists", kind="docx"), None),
    (chk("docx_words_between", min=100, max=300), None),
    (chk("xlsx_columns", names=["Item", "EUR"]), None),
    (chk("verdict_present"), None),
    (chk("made_up"), "unknown check"),
    (chk("file_exists", kind="exe"), "kind must be one of"),
    (chk("docx_pages_max"), "missing parameter n"),
    (chk("docx_pages_max", n=-1), "whole number"),
    (chk("docx_words_between", min=300, max=100), "min is above max"),
    (chk("output_mentions", terms=[]), "non-empty list"),
    (chk("csv_rows_min", n=3, sheet="x"), "unknown parameter"),
    (chk("figures_dated", covers=["requirement 1"]), "covers must be"),
])
def test_check_problem(c, problem):
    got = C.check_problem(c)
    assert (got is None) if problem is None else (problem in got), got


def test_the_schema_vocabulary_matches_the_evaluator():
    assert set(CheckOut.model_fields["name"].annotation.__args__) == set(C.VOCAB)


def test_legacy_plan_reads_an_optional_checks_line():
    text = ("1. [Analyst]: Compute\n   covers: R1\n   output: table: costs\n   done_when: costs given\n"
            '   checks: [{"name": "output_has_table", "params": {"min_rows": 3}, "covers": ["R1"]}]\n'
            "2. [Writer]: Write\n   covers: R1\n   checks: not json\n")
    (_, _, f1), (_, _, f2) = parse_plan_d24(text)
    assert f1["checks"] == [chk("output_has_table", min_rows=3)]
    assert f2["checks"][0]["name"] == "unparsed" and C.check_problem(f2["checks"][0])


def test_a_plan_without_checks_derives_them_from_the_markers():
    assert C.derived_checks("table: x; list: y", ["R2"]) == [
        {"name": "output_has_table", "params": {"min_rows": 1}, "covers": ["R2"], "derived": True},
        {"name": "output_has_list", "params": {"min_items": 2}, "covers": ["R2"], "derived": True}]
    assert C.derived_checks("memo: a report", ["R1"]) == []


def test_planout_carries_checks_into_the_legacy_text():
    p = PlanOut(requirements=[], givens=[], open_questions=[], roles=[], capability_requests=[], risks=[],
                plan=[{"number": 1, "roles": ["A"], "title": "t", "kind": "work", "covers": ["R1"], "depends_on": [],
                       "do": "d", "output": "table: x", "done_when": "w",
                       "checks": [chk("output_has_table", min_rows=2)]}])
    plan = plan_markdown(p).split("## Execution Plan:\n")[1].split("\n\n##")[0]
    [(_, _, f)] = parse_plan_d24(plan)
    assert f["checks"] == [chk("output_has_table", min_rows=2)]


# ---- the evaluators ------------------------------------------------------------------------------------------------
@pytest.fixture
def files(tmp_path):
    import docx
    import openpyxl
    from PIL import Image
    d = docx.Document()
    d.add_paragraph(" ".join(["word"] * 1200))
    t = d.add_table(rows=4, cols=2)
    t.cell(0, 0).text = "a"
    d.save(tmp_path / "memo.docx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Item", "EUR"])
    for i in range(5):
        ws.append([f"x{i}", i])
    ws.append(["Total", "=SUM(B2:B6)"])
    wb.save(tmp_path / "costs.xlsx")
    (tmp_path / "data.csv").write_text("city,km\nA,1\nB,2\n", encoding="utf-8")
    Image.new("RGB", (640, 480)).save(tmp_path / "chart.png")
    (tmp_path / "email_NOT_SENT.md").write_text("Draft. NOT SENT.", encoding="utf-8")
    return [str(p) for p in tmp_path.iterdir()]


@pytest.mark.parametrize("c,ok", [
    (chk("file_exists", kind="docx"), True),
    (chk("file_exists", kind="pdf"), False),
    (chk("file_exists", kind="xlsx", name="costs.xlsx"), True),
    (chk("docx_pages_max", n=2), False),                 # 1200 words, no page breaks: 3 pages estimated
    (chk("docx_pages_max", n=3), True),
    (chk("docx_words_between", min=1000, max=1500), True),
    (chk("docx_has_table", min_rows=4), True),
    (chk("docx_has_table", min_rows=5), False),
    (chk("xlsx_rows_min", n=6), True),
    (chk("xlsx_rows_min", n=7), False),
    (chk("xlsx_columns", names=["item", "EUR"]), True),
    (chk("xlsx_columns", names=["VAT"]), False),
    (chk("xlsx_has_formulas"), True),
    (chk("csv_rows_min", n=2), True),
    (chk("csv_columns", names=["city", "km", "litres"]), False),
    (chk("image_exists", kind="png", min_px=400), True),
    (chk("image_exists", kind="png", min_px=500), False),
    (chk("image_exists", kind="svg"), False),
])
def test_file_checks(files, c, ok):
    r = C.run_check(c, {"text": "", "files": files})
    assert r["pass"] is ok, r["detail"]
    assert r["name"] == f"contract:{c['name']}" and r["covers"] == ["R1"]


def test_docx_pages_say_estimated(files):
    assert "estimated" in C.run_check(chk("docx_pages_max", n=9), {"files": files})["detail"]


TEXT = ("Verdict: PASS\n\n| a | b |\n|---|---|\n| 1 | 2 |\n| 3 | 4 |\n\n- one\n- two\n- three\n\nDiesel costs 1.80 EUR.\n\n"
        "## Limitations\nThe email was NOT SENT; it is saved as email_NOT_SENT.md.")


@pytest.mark.parametrize("c,ctx,ok", [
    (chk("output_has_table", min_rows=2), {}, True),
    (chk("output_has_table", min_rows=3), {}, False),
    (chk("output_has_list", min_items=3), {}, True),
    (chk("output_has_list", min_items=4), {}, False),
    (chk("output_mentions", terms=["DIESEL", "eur"]), {}, True),
    (chk("output_mentions", terms=["petrol"]), {}, False),
    (chk("cites_every_figure"), {"untagged": 0}, True),
    (chk("cites_every_figure"), {"untagged": 2}, False),
    (chk("figures_dated"), {"undated": []}, True),
    (chk("figures_dated"), {"undated": [{"figure": "1.80"}]}, False),
    (chk("verdict_present"), {"verdict": "PASS"}, True),
    (chk("verdict_present"), {"verdict": None}, False),
])
def test_text_checks(c, ctx, ok):
    assert C.run_check(c, {"text": TEXT, **ctx})["pass"] is ok


def test_not_sent_fallback(files):
    assert C.run_check(chk("not_sent_fallback"), {"text": TEXT, "files": files})["pass"]
    assert not C.run_check(chk("not_sent_fallback"), {"text": "Sent.", "files": files})["pass"]
    assert not C.run_check(chk("not_sent_fallback"), {"text": TEXT, "files": []})["pass"]


def test_an_invalid_check_fails_without_crashing(files):
    r = C.run_check(chk("docx_pages_max"), {"files": files})
    assert r["pass"] is False and r["detail"].startswith("invalid check")
    broken = files[0].rsplit("/", 1)[0] + "/bad.docx"
    open(broken, "w").write("not a docx")
    assert C.run_check(chk("docx_words_between", min=0, max=9), {"files": [broken]})["pass"] is False


# ---- Box 2: the HARD checks ----------------------------------------------------------------------------------------
def step(n, roles, covers, checks):
    return DraftPlanStep(index=n - 1, agent_names=roles, text=f"[{roles[0]}]: s{n}", covers=covers, checks=checks)


def test_plan_contract_quality():
    tools = {"Analyst": ["calc"], "Writer": ["local:Write"]}
    good_plan = [step(1, ["Analyst"], ["R1"], [chk("output_has_table", ["R1"])]),
                 step(2, ["Writer"], ["R1", "R2"], [chk("file_exists", ["R2"], kind="docx")])]
    q = C.plan_contract_quality(["R1", "R2"], good_plan, tools)
    assert all(q[k]["ok"] for k in q), q
    bad = [step(1, ["Analyst"], ["R1"], [chk("file_exists", ["R1"], kind="docx"), chk("nope", ["R1"])] +
                [chk("output_mentions", ["R1"], terms=["x"])] * 5)]
    q = C.plan_contract_quality(["R1", "R2"], bad, tools)
    assert q["requirements_checked"] == {"ok": False, "unchecked": ["R2"]}
    assert not q["checks_valid"]["ok"] and "unknown check" in q["checks_valid"]["invalid"][0]["problem"]
    assert q["checks_reachable"]["too_many"] == [{"step": 1, "checks": 7}]
    assert q["checks_reachable"]["no_file_tool"] == [{"step": 1, "checks": ["file_exists"]}]
    lines = gate_suggestions({"checks": q, "hard_failed": ["requirements_checked", "checks_valid", "checks_reachable"]})
    assert "R2" in lines and "unknown check" in lines and "at most 6" in lines and "file-writing tool" in lines


def test_a_plan_with_no_checks_is_not_judged():
    q = C.plan_contract_quality(["R1"], [step(1, ["A"], ["R1"], [])], {})
    assert all(v["ok"] is None for v in q.values())


# ---- Box 3 ----------------------------------------------------------------------------------------------------------
def run_with(task, envelope, trace, tools, checks: dict[int, list]):
    llm, cfg = plan_team(task, envelope, trace, plan_worker=lambda m, s: good(step_no(m)))
    for s in cfg.plan:
        s.checks = checks.get(s.index + 1, [])
    ep = Interpreter(llm, tools, trace).run(cfg, task, seed=0)
    return {m["step"]: m for m in ep.steps}, trace


def test_box3_runs_a_steps_checks_and_records_them(task, envelope, trace, tools):
    metas, tr = run_with(task, envelope, trace, tools, {1: [chk("output_has_table", min_rows=1)],
                                                        2: [chk("output_mentions", terms=["zebra"])]})
    assert metas[1]["contract_checks"][0]["pass"] and metas[1]["status"] == "done"
    m2 = metas[2]
    assert m2["contract_checks"] == [c for c in m2["checks"] if c.get("source") == "contract"]
    assert not m2["contract_checks"][0]["pass"] and "zebra" in m2["contract_checks"][0]["detail"]
    assert m2["status"] == "incomplete" and "contract:output_mentions" in m2["status_reason"]
    assert m2["retried"]                                         # the D61 check-retry turn ran first
    assert any(e["amoeba.step"] == 2 for e in tr.events("contract_checks"))
    assert metas[3]["contract_checks"] == []                     # no checks: nothing new


def test_a_failed_check_feeds_box4s_checks_cause(task, envelope, trace, tools):
    from amoeba.adapt.stuck import diagnose, step_signals
    metas, _ = run_with(task, envelope, trace, tools, {2: [chk("output_mentions", terms=["zebra"])]})
    d = diagnose(step_signals(metas[2], plan={}))
    assert d["cause"] == "checks"
    assert any("contract:output_mentions" in w for w in d.get("why") or d.get("evidence") or [str(d)])


# ---- the final requirement check -----------------------------------------------------------------------------------
ANSWER = "## Answer\nThe weekly diesel cost is 97.20 EUR.\n\n## Limitations\nThe route map was not made."


def test_met_needs_every_covering_check_passed():
    passed = [{"check": "output_mentions", "covers": ["R1"], "pass": True}]
    failed = [{"check": "file_exists", "covers": ["R1"], "pass": False}]
    assert requirement_check("R1", "weekly diesel cost", ANSWER, None, passed)["status"] == "met"
    r = requirement_check("R1", "weekly diesel cost", ANSWER, None, failed)
    assert r["status"] == "missing" and r["status_legacy"] == "met" and r["checks_failed"] == ["file_exists"]


def test_terms_only_under_limitations_are_reported_missing():
    r = requirement_check("R2", "a route map", ANSWER, None, [])
    assert r["status"] == "reported_missing" and r["status_legacy"] == "met"
    fc = final_check(ANSWER + " " * 400, {"R1": "weekly diesel cost", "R2": "a route map"}, {}, None, checks=[])
    assert fc["unmet"] == ["R2"] and fc["reported_missing"] == ["R2"] and fc["met_legacy"] == ["R1", "R2"]
    assert not any("requirements not met" in x for x in fc["core_missing"])   # said honestly: not no_deliverable
    from amoeba.task.deliverables import final_findings
    assert final_findings(fc) == []                               # nothing to ask the refine turn for


def test_result_json_keeps_met_legacy():
    fc = final_check("Plain answer about diesel cost." * 20, {"R1": "diesel cost"}, {}, None,
                     checks=[{"check": "figures_dated", "covers": ["R1"], "pass": False}])
    assert fc["met_legacy"] == ["R1"] and fc["unmet"] == ["R1"]
    json.dumps(fc)


def test_helpers_are_told_what_will_be_checked():
    from amoeba.config.schema import PlanStep
    from amoeba.interp.plan_runner import step_detail
    s = PlanStep(index=0, agent_ids=["a"], text="[A]: t", done_when="w")
    assert "checked by plain code" not in step_detail(s)                 # no checks: the prompt is unchanged
    s.checks = [chk("docx_pages_max", n=2), chk("verdict_present")]
    assert "checked by plain code when you finish: docx_pages_max (n=2); verdict_present" in step_detail(s)
