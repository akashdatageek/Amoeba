"""D171 — real task scorers (amoeba/eval/sim_score.py SCORERS), all part of the frozen judge:
answer_key (numbers with a tolerance, exact strings), hidden_tests (a pytest folder the team never sees, run next to a
copy of the delivered code), files (plain-code checks on delivered files) and rubric (a fixed grader prompt, a fixed
model and temperature, 3 calls, the median). A task may weight several scorers. Everything a scorer reads is in the
judge hash. Offline: the rubric grader is a recorded mock."""
import json
import shutil
from pathlib import Path

import pytest

import amoeba.eval.sim_score as S
from amoeba.eval.judge import frozen_paths, judge_hash
from amoeba.llm.client import MockLLMClient
from amoeba.task.models import Task
from tests.test_frozen_judge_d161 import copy_judge

ROOT = Path(__file__).resolve().parents[1]


def run_dir(tmp_path, answer="", files=None):
    d = tmp_path / "run"
    (d / "workspace").mkdir(parents=True, exist_ok=True)
    (d / "result.json").write_text(json.dumps({"answer": answer}))
    for name, data in (files or {}).items():
        p = d / "workspace" / name
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_bytes(data if isinstance(data, bytes) else data.encode())
    return d


def keys(tmp_path, task_id, key):
    folder = tmp_path / "tasks" / "demo_keys"
    folder.mkdir(parents=True, exist_ok=True)
    (folder / f"{task_id}.json").write_text(json.dumps(key))
    return str(tmp_path / "tasks" / "demo")                 # the judge_set prefix


def task(tmp_path, tid, sim):
    return Task(id=tid, prompt="Work out the figure and name the database.", sim=sim)


def score(t, d, seed=0):
    return S.score_task(d, t, seed)


def test_answer_key_numbers_with_tolerance_and_exact_strings(tmp_path):
    k = keys(tmp_path, "t1", {"numbers": [{"name": "total", "value": 396, "tol": 0.5},
                                          {"name": "share", "value": 0.25, "rel_tol": 0.02}],
                              "strings": ["PostgreSQL"]})
    t = task(tmp_path, "t1", {"scorers": [{"name": "answer_key"}], "judge_set": k})
    full = run_dir(tmp_path, "The total is 396.2 and the share is 0.251; pick PostgreSQL.")
    assert score(t, full)["score"] == 1.0 and score(t, full) == score(t, full)          # deterministic
    half = run_dir(tmp_path / "b", "The total is 390; the share is 0.25. Pick MongoDB.")
    got = score(t, half)
    assert abs(got["score"] - 1 / 3) < 1e-6
    assert [x["ok"] for x in got["parts"][0]["detail"]] == [False, True, False]


def test_hidden_tests_run_next_to_a_copy_of_the_code_and_score_the_share_passed(tmp_path):
    hidden = tmp_path / "tasks" / "demo_hidden" / "t2"
    hidden.mkdir(parents=True)
    (hidden / "test_tool.py").write_text("from tool import add\n\ndef test_add():\n    assert add(2, 3) == 5\n\n"
                                         "def test_neg():\n    assert add(-1, -1) == -3\n")
    t = task(tmp_path, "t2", {"scorers": [{"name": "hidden_tests"}], "judge_set": str(tmp_path / "tasks" / "demo")})
    d = run_dir(tmp_path, "done", {"tool.py": "def add(a, b):\n    return a + b\n"})
    got = score(t, d)
    assert got["score"] == 0.5 and score(t, d)["score"] == 0.5
    assert sorted(p.name for p in (d / "workspace").iterdir()) == ["tool.py"]    # never copied into the workspace


def test_files_checks_on_delivered_files(tmp_path):
    import openpyxl
    from io import BytesIO
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["A2"], ws["A3"] = 2, 3, "=A1*A2"
    buf = BytesIO()
    wb.save(buf)
    k = keys(tmp_path, "t3", {"files": [{"check": "formulas_live", "path": "*.xlsx"},
                                        {"check": "rows", "path": "out.csv", "min": 3, "max": 3},
                                        {"check": "exists", "path": "*.png"},
                                        {"check": "pages", "path": "*.pdf", "min": 1}]})
    t = task(tmp_path, "t3", {"scorers": [{"name": "files"}], "judge_set": k})
    d = run_dir(tmp_path, "see files", {"book.xlsx": buf.getvalue(), "out.csv": "a,b\n1,2\n3,4\n5,6\n",
                                        "chart.png": b"\x89PNG\r\n\x1a\nxx"})
    got = score(t, d)
    assert got["score"] == 0.75 and score(t, d) == got                               # no pdf: 3 of 4
    assert [x["ok"] for x in got["parts"][0]["detail"]] == [True, True, True, False]


def grader(replies):
    return MockLLMClient(script={"rubric_grader": list(replies)}, model="gemma-4-31b-it")


def test_rubric_is_the_median_of_three_fixed_grader_calls(tmp_path, monkeypatch):
    k = keys(tmp_path, "t4", {"rubric": "1 point: names the trade-off. 1 point: states what is unknown."})
    t = task(tmp_path, "t4", {"scorers": [{"name": "rubric"}], "judge_set": k})
    d = run_dir(tmp_path, "Postgres, because joins; the load is unknown.")
    for _ in range(2):                                    # the same recorded replies give the same median
        g = grader(["SCORE: 0.4", "SCORE: 0.9", "SCORE: 0.6"])
        monkeypatch.setattr(S, "grader_client", lambda cfg: g)
        got = score(t, d)
        assert got["score"] == 0.6 and len(g.calls) == 3
    prompt = g.calls[0]["messages"][-1]["content"]
    assert t.prompt in prompt and "names the trade-off" in prompt and "the load is unknown" in prompt
    assert all(c["temperature"] == 0.0 for c in g.calls) if "temperature" in g.calls[0] else True
    assert got["parts"][0]["grader"]["model"] == "gemma-4-31b" and got["parts"][0]["grader"]["calls"] == 3


def test_several_scorers_with_weights(tmp_path):
    k = keys(tmp_path, "t5", {"numbers": [{"value": 10, "tol": 0}], "files": [{"check": "exists", "path": "x.md"}]})
    t = task(tmp_path, "t5", {"scorers": [{"name": "answer_key", "weight": 3}, {"name": "files", "weight": 1}],
                              "judge_set": k})
    assert score(t, run_dir(tmp_path, "10"))["score"] == 0.75


def test_a_changed_key_or_hidden_test_changes_the_judge_hash(tmp_path):
    root = copy_judge(tmp_path)
    for rel in ("tasks/demo_keys/a.json", "tasks/demo_hidden/a/test_a.py"):
        (root / rel).parent.mkdir(parents=True, exist_ok=True)
        (root / rel).write_text("1")
    before = judge_hash(root)["sha256"]
    (root / "tasks/demo_keys/a.json").write_text("2")
    mid = judge_hash(root)["sha256"]
    (root / "tasks/demo_hidden/a/test_a.py").write_text("2")
    assert len({before, mid, judge_hash(root)["sha256"]}) == 3
    paths = frozen_paths()
    for p in ("tasks/*_keys", "tasks/*_hidden", "amoeba/config/prompts/rubric_grader.txt", "amoeba/eval"):
        assert p in paths


def test_a_task_may_not_be_given_its_own_key_or_hidden_tests(tmp_path):
    k = keys(tmp_path, "t6", {"numbers": [{"value": 1, "tol": 0}]})
    t = Task(id="t6", prompt="x", files=[str(Path(k + "_keys") / "t6.json")],
             sim={"scorers": [{"name": "answer_key"}], "judge_set": k})
    with pytest.raises(ValueError, match="key"):
        S.check_no_leak(t)
