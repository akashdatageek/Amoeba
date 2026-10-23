import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
sol=json.load(open('solution.json'))
oa=sol['optimal_assignment']
F='Arial'
base=Font(name=F,size=10); bold=Font(name=F,size=10,bold=True); blue=Font(name=F,size=10,color='0000FF')
hdr=PatternFill('solid',fgColor='1F3864'); hf=Font(name=F,size=10,bold=True,color='FFFFFF')
inp=PatternFill('solid',fgColor='FFF2CC')
USD='$#,##0.00'; INT='#,##0'
wb=Workbook()
def head(ws,row,names):
    for i,n in enumerate(names,1):
        c=ws.cell(row,i,n); c.font=hf; c.fill=hdr; c.alignment=Alignment(horizontal='center',wrap_text=True)
loads=[("L1",18000,220),("L2",12500,340),("L3",22000,150),("L4",9000,410),("L5",15500,95),("L6",7000,280),("L7",20500,190),("L8",11000,365),("L9",14000,120),("L10",8500,255)]
trucks=[("A",45000,1100,350,2.10),("B",45000,1100,350,2.10),("C",30000,1100,500,1.85),("D",26000,900,275,2.45)]
# Inputs
wi=wb.active; wi.title='Inputs'
wi['A1']='Inputs — carrier load/truck data (source: task statement; blue = hard-coded input)'; wi['A1'].font=Font(name=F,size=12,bold=True)
head(wi,4,['Load ID','Weight (lb)','Miles'])
for i,(a,w,m) in enumerate(loads):
    r=5+i; wi.cell(r,1,a); wi.cell(r,2,w).number_format=INT; wi.cell(r,3,m).number_format=INT
head(wi,18,['Truck','Capacity (lb)','Mileage limit','Fixed $/day','Rate $/mile'])
for i,(t,c,l,f,rt) in enumerate(trucks):
    r=19+i; wi.cell(r,1,t); wi.cell(r,2,c).number_format=INT; wi.cell(r,3,l).number_format=INT
    wi.cell(r,4,f).number_format=USD; wi.cell(r,5,rt).number_format=USD
wi['A3']='Load table'; wi['A17']='Truck table'; wi['A3'].font=bold; wi['A17'].font=bold
for row in wi.iter_rows(min_row=5,max_row=22):
    for c in row:
        if c.value is not None: c.font=blue
for a in ('A3','A17'): wi[a].font=bold
for col,w in zip('ABCDE',(12,16,16,14,14)): wi.column_dimensions[col].width=w
# Assignment
wa=wb.create_sheet('Assignment')
head(wa,1,['Load','Truck','Weight','Miles','','A','B','C','D'])
wa['F1']='Helper (list build) →'; wa['F1'].fill=hdr; wa['F1'].font=hf
# helper headers G..J = truck letters via formulas
for j,col in enumerate('GHIJ'):
    wa[f'{col}1']=f'=Inputs!A{19+j}'; wa[f'{col}1'].fill=hdr; wa[f'{col}1'].font=hf; wa[f'{col}1'].alignment=Alignment(horizontal='center')
for ch in ('E','F'): pass
wa['E1']=None; 
asg={}
for t,d in oa.items():
    for l in d['loads']: asg[l]=t
for i in range(10):
    r=2+i
    wa.cell(r,1,f'=Inputs!A{5+i}')
    c=wa.cell(r,2,asg[loads[i][0]]); c.fill=inp; c.font=blue; c.alignment=Alignment(horizontal='center')
    wa.cell(r,3,f'=INDEX(Inputs!$B$5:$B$14,MATCH($A{r},Inputs!$A$5:$A$14,0))').number_format=INT
    wa.cell(r,4,f'=INDEX(Inputs!$C$5:$C$14,MATCH($A{r},Inputs!$A$5:$A$14,0))').number_format=INT
    for col in 'GHIJ':
        wa[f'{col}{r}']=f'=IF($B{r}={col}$1,", "&$A{r},"")'
    for col in 'ACD': wa[f'{col}{r}'].font=base
dv=DataValidation(type='list',formula1='"A,B,C,D"',allow_blank=False); wa.add_data_validation(dv); dv.add('B2:B11')
wa['A13']='Edit the yellow Truck cells (B2:B11) to try other plans; Weight and Miles are looked up from Inputs by formula. Columns G:J are helpers for the load lists on the Cost sheet.'
wa['A13'].font=Font(name=F,size=9,italic=True)
for col,w in zip('ABCDEFGHIJ',(10,10,12,10,3,20,12,12,12,12)): wa.column_dimensions[col].width=w
# Cost
wc=wb.create_sheet('Cost')
head(wc,1,['Truck','Weight','Miles','Loads count','Used','Weight OK','Miles OK','Fixed','Variable','Truck cost','Loads'])
for j in range(4):
    r=2+j; ir=19+j; hc='GHIJ'[j]
    wc[f'A{r}']=f'=Inputs!A{ir}'
    wc[f'B{r}']=f'=SUMIF(Assignment!$B$2:$B$11,$A{r},Assignment!$C$2:$C$11)'
    wc[f'C{r}']=f'=SUMIF(Assignment!$B$2:$B$11,$A{r},Assignment!$D$2:$D$11)'
    wc[f'D{r}']=f'=COUNTIF(Assignment!$B$2:$B$11,$A{r})'
    wc[f'E{r}']=f'=IF(D{r}>0,"Yes","No")'
    wc[f'F{r}']=f'=IF(B{r}<=Inputs!B{ir},"OK","OVER")'
    wc[f'G{r}']=f'=IF(C{r}<=Inputs!C{ir},"OK","OVER")'
    wc[f'H{r}']=f'=IF(E{r}="Yes",Inputs!D{ir},0)'
    wc[f'I{r}']=f'=Inputs!E{ir}*C{r}'
    wc[f'J{r}']=f'=H{r}+I{r}'
    X='&'.join(f'Assignment!{hc}{k}' for k in range(2,12))
    wc[f'K{r}']=f'=MID({X},3,200)'
    for col in 'BCD': wc[f'{col}{r}'].number_format=INT
    for col in 'HIJ': wc[f'{col}{r}'].number_format=USD
wc['A6']='Total'
for col in 'BCD': wc[f'{col}6']=f'=SUM({col}2:{col}5)'; wc[f'{col}6'].number_format=INT
for col in 'HIJ': wc[f'{col}6']=f'=SUM({col}2:{col}5)'; wc[f'{col}6'].number_format=USD
wc['F6']='=IF(COUNTIF(F2:F5,"OK")=4,"OK","CHECK")'; wc['G6']='=IF(COUNTIF(G2:G5,"OK")=4,"OK","CHECK")'
for c in wc[6]: c.font=bold
wc['A8']='TOTAL COST'; wc['B8']='=J6'; wc['B8'].number_format=USD
wc['A8'].font=Font(name=F,size=12,bold=True); wc['B8'].font=Font(name=F,size=12,bold=True)
wc['A10']='Checks'; wc['A10'].font=bold
chk=[('Loads assigned (COUNTA of assignments = 10)','=COUNTA(Assignment!B2:B11)','=IF(B11=COUNTA(Inputs!A5:A14),"OK","CHECK")'),
('Assignments with a valid truck letter','=SUMPRODUCT(COUNTIF(Inputs!A19:A22,Assignment!B2:B11))','=IF(B12=COUNTA(Inputs!A5:A14),"OK","CHECK")'),
('Sum of per-truck load counts (each load exactly once)','=D6','=IF(B13=COUNTA(Inputs!A5:A14),"OK","CHECK")'),
('Total weight assigned = total weight of loads','=B6','=IF(B14=SUM(Inputs!B5:B14),"OK","CHECK")'),
('Weight limits, all trucks','=COUNTIF(F2:F5,"OK")','=IF(B15=4,"OK","CHECK")'),
('Mileage limits, all trucks','=COUNTIF(G2:G5,"OK")','=IF(B16=4,"OK","CHECK")')]
for i,(l,f,o) in enumerate(chk):
    r=11+i; wc[f'A{r}']=l; wc[f'B{r}']=f; wc[f'C{r}']=o
wc['A17']='ALL CHECKS'; wc['C17']='=IF(COUNTIF(C11:C16,"OK")=6,"OK","CHECK")'
wc['A17'].font=bold; wc['C17'].font=bold
wc['A18']='Total cost equals solver optimum (Optimality!B5)'; wc['C18']='=IF(ABS(B8-Optimality!B5)<0.005,"OK","DIFFERS")'
wb_cost=wc
for col,w in zip('ABCDEFGHIJK',(52,14,12,12,8,11,11,12,12,13,22)): wc.column_dimensions[col].width=w
for row in wc.iter_rows():
    for c in row:
        if c.font==Font() or c.font.name!=F:
            c.font=Font(name=F,size=c.font.size if c.font.size and c.font.size!=11 else 10,bold=c.font.bold,color=c.font.color)
# Optimality
wo=wb.create_sheet('Optimality')
wo['A1']='Optimality evidence — SOLVER OUTPUT (hard-coded text/values from solution.json; not recomputed by this workbook)'; wo['A1'].font=Font(name=F,size=12,bold=True)
top=sol['top5_distinct_costs']
rows=[('Method','Exhaustive enumeration of all 4^10 = 1,048,576 assignments; no pruning'),
('Feasible assignments (weight & mileage OK)',int(sol['feasible_count'])),
('Optimal cost (solver)',float(sol['total_cost'])),
('Number of optimal assignments',int(sol['optimal_count'])),
('Next-best distinct cost (solver)',float(top[1]['cost'])),
('Assignments at next-best cost',int(top[1]['num_assignments'])),
]
# put optimal cost in B5
labels=[('A3','Method'),]
r=3
for l,v in rows:
    wo.cell(r,1,l).font=bold; c=wo.cell(r,2,v); c.font=blue; r+=1
# rows: 3 method,4 feasible,5 optimal cost,6 num optimal,7 next best,8 count next
wo['B5'].number_format=USD; wo['B7'].number_format=USD; wo['B4'].number_format=INT
wo['A9']='Margin to next-best (formula)'; wo['B9']='=B7-B5'; wo['B9'].number_format=USD; wo['A9'].font=bold
wo['A10']='Workbook total cost (formula)'; wo['B10']='=Cost!B8'; wo['B10'].number_format=USD; wo['A10'].font=bold
wo['A11']='Workbook total = solver optimum?'; wo['B11']='=IF(ABS(B10-B5)<0.005,"OK","DIFFERS")'; wo['A11'].font=bold
wo['A13']='Distinct cost levels (solver top 5)'; wo['A13'].font=bold
wo['A14']='Cost'; wo['B14']='Assignments'
for c in (wo['A14'],wo['B14']): c.font=hf; c.fill=hdr
for i,t in enumerate(top):
    wo.cell(15+i,1,float(t['cost'])).number_format=USD; wo.cell(15+i,2,int(t['num_assignments']))
    wo.cell(15+i,1).font=blue; wo.cell(15+i,2).font=blue
wo['A21']='All four trucks are needed (inputs-based formulas)'; wo['A21'].font=bold
wo['A22']='Total load weight (lb)'; wo['B22']='=SUM(Inputs!B5:B14)'
wo['A23']='Three largest capacities (lb)'; wo['B23']='=LARGE(Inputs!B19:B22,1)+LARGE(Inputs!B19:B22,2)+LARGE(Inputs!B19:B22,3)'
wo['A24']='Three trucks enough?'; wo['B24']='=IF(B23>=B22,"Yes","No – all 4 trucks required")'
wo['A25']='Fixed cost of all four trucks'; wo['B25']='=SUM(Inputs!D19:D22)'; wo['B25'].number_format=USD
for a in ('B22','B23'): wo[a].number_format=INT
wo['A27']='Solver notes (text)'; wo['A27'].font=bold
notes=['Total load weight 138,000 lb > 120,000 lb (three largest capacities), so every feasible plan uses all four trucks; fixed cost is therefore always $1,475.00 and only variable cost is optimised.',
'Trucks A and B are identical, so the 2 optimal assignments are a mirror pair (A and B load sets swapped); C = {L4, L6, L8}, D = {L3} in both. Up to the A/B swap the optimum is unique.',
'Optimal plan: A = L1, L2, L9; B = L5, L7, L10; C = L4, L6, L8; D = L3. Source: solution.json (exact Fraction arithmetic).']
for i,n in enumerate(notes):
    c=wo.cell(28+i,1,n); wo.merge_cells(start_row=28+i,start_column=1,end_row=28+i,end_column=2)
    c.alignment=Alignment(wrap_text=True,vertical='top'); wo.row_dimensions[28+i].height=45
wo.column_dimensions['A'].width=62; wo.column_dimensions['B'].width=48
for ws in (wo,):
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and c.font.name!=F: c.font=Font(name=F,size=10)
wo['B3'].alignment=Alignment(wrap_text=True)
for ws in (wi,wa):
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and c.font.name!=F: c.font=base
wb.move_sheet('Assignment', offset=-1)
wb.save('truck_assignment_plan.xlsx')
