import openpyxl

def audit():
    # Load with formulas
    wb_f = openpyxl.load_workbook('truck_assignment_plan.xlsx', data_only=False)
    # Load with values
    wb_v = openpyxl.load_workbook('truck_assignment_plan.xlsx', data_only=True)
    
    print("--- Sheet Names ---")
    print(wb_f.sheetnames)
    
    # 1. Audit Inputs Sheet
    print("\n--- Inputs Sheet (Values) ---")
    ins_v = wb_v['Inputs']
    for r in range(1, 23):
        print(f"Row {r}: {ins_v[f'A{r}'].value}, {ins_v[f'B{r}'].value}, {ins_v[f'C{r}'].value}")

    # 2. Audit Assignment Sheet
    print("\n--- Assignment Sheet (Values) ---")
    ass_v = wb_v['Assignment']
    for r in range(1, 12):
        print(f"Row {r}: {ass_v[f'A{r}'].value}, {ass_v[f'B{r}'].value}, {ass_v[f'C{r}'].value}, {ass_v[f'D{r}'].value}")

    # 3. Audit Cost Sheet Formulas
    print("\n--- Cost Sheet Formulas ---")
    cost_f = wb_f['Cost']
    for r in range(2, 6):
        print(f"Row {r} (Truck {cost_f[f'A{r}'].value}):")
        print(f"  Weight (B): {cost_f[f'B{r}'].value}")
        print(f"  Miles (C): {cost_f[f'C{r}'].value}")
        print(f"  Fixed (H): {cost_f[f'H{r}'].value}")
        print(f"  Variable (I): {cost_f[f'I{r}'].value}")
        print(f"  Total (J): {cost_f[f'J{r}'].value}")
    
    print(f"Total Cost (J6): {cost_f['J6'].value}")
    print(f"B8: {cost_f['B8'].value}")

    # 4. Audit Cost Sheet Values
    print("\n--- Cost Sheet Values ---")
    cost_v = wb_v['Cost']
    for r in range(2, 6):
        print(f"Row {r}: {cost_v[f'J{r}'].value}")
    print(f"Total (J6): {cost_v['J6'].value}")

audit()
