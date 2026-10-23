import csv
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.section import WD_ORIENT
# CSV
rows=[]
with open('probe_averages.csv',newline='') as f:
    for r in csv.DictReader(f):
        rows.append(r)
with open('probe_row_map_used.csv','w',newline='') as f:
    w=csv.writer(f); w.writerow(['probe_id','row_id','avg_vwc_pct','below_20'])
    for r in rows:
        w.writerow([r['probe_id'],'MAPPING REQUIRED',r['avg_vwc_pct'],r['below_20']])
d=Document()
s=d.sections[0]; s.orientation=WD_ORIENT.LANDSCAPE; s.page_width,s.page_height=s.page_height,s.page_width
d.add_heading('Irrigation Order - Field 7',0)
p=d.add_paragraph(); r=p.add_run('DRAFT - NOT ISSUED'); r.bold=True; r.font.size=Pt(20); r.font.color.rgb=RGBColor(0xC0,0,0)
p=d.add_paragraph(); r=p.add_run('BANNER: Soil-moisture data NOT PULLED and probe-to-row mapping MISSING. No probe has an average VWC, so no row can be named for irrigation. This order was NOT SENT; it is saved as a file only.'); r.bold=True
d.add_heading('Header block',1)
for k,v in [
 ('Field','Field 7'),
 ('Prepared','Friday 23 October 2026 (America/Chicago)'),
 ('Data window (UTC)','Intended: last 24 h, [pull time - 24 h, pull time] in UTC. NOT PULLED, so no actual UTC start/end exists.'),
 ('Data source','ChirpStack v4 at 10.20.0.5 (REST/gRPC), Field 7 application or device group; VWC % as reported in the probe payload, no conversion'),
 ('Data status','NOT PULLED - no API token, no network from the sandbox (step 1). Scripts verified on synthetic data only (step 2 verdict PASS; covers scripts, not farm data).'),
 ('Mapping status','MISSING - no probe-to-row layout or probe-location list supplied'),
 ('Threshold','Flag a probe when its 24 h mean VWC is strictly below 20.0 %'),
 ('Delivery','NOT SENT - saved as a file only')]:
    p=d.add_paragraph(); p.add_run(k+': ').bold=True; p.add_run(v)
d.add_heading('Irrigation table (placeholders - nothing to irrigate can be named yet)',1)
hdr=['Priority','Row','Probes','Avg VWC %','Action','Volume / Duration / Start time']
t=d.add_table(rows=1,cols=6); t.style='Table Grid'
for i,h in enumerate(hdr):
    c=t.rows[0].cells[i]; c.text=''; c.paragraphs[0].add_run(h).bold=True
vals=['-','MAPPING REQUIRED','None identified: 12 probes, all NOT PULLED (IDs are placeholders PROBE_ID_UNKNOWN_01 to _12)','NOT PULLED (blank)','Irrigate - only for rows of probes below 20 %, once data and mapping exist','TO BE CONFIRMED by farm manager']
c=t.add_row().cells
for i,v in enumerate(vals): c[i].text=v
d.add_heading('Notes',1)
for n in [
 'No application rates, volumes, durations or start times are given. They are TO BE CONFIRMED by the farm manager from farm-supplied parameters.',
 'When real averages exist: take probes with below_20 TRUE, map each to its row, list each row once, and order rows driest first (lowest average VWC of its flagged probes; the farm manager should confirm how a row with several flagged probes is ranked). Unmapped probes stay under row MAPPING REQUIRED.',
 'LOW COVERAGE: any probe with fewer than 75 % of the 24 expected hourly readings is marked LOW COVERAGE by the analysis script; its average is less reliable and must be checked before it drives an irrigation decision. A probe with no data has no average and is not flagged. Currently all 12 probes are NOT PULLED.',
 'Missing inputs: CHIRPSTACK_TOKEN; Field 7 application or device-group ID and REST port (script default 8090); confirmation that the payload key is vwc; the probe-to-row layout; the real probe IDs and device count.',
 'To complete on site: export CHIRPSTACK_TOKEN, run fetch_chirpstack.py then analyse_vwc.py, add a probe_id,row_id mapping file, then re-issue this order.',
 'The fetch script REST paths are untested against the live server (step 2).',
 'Full probe list: probe_row_map_used.csv (all rows MAPPING REQUIRED, averages blank).']:
    d.add_paragraph(n,style='List Bullet')
d.add_heading('Sign-off',1)
for l in ['Prepared by (Irrigation Planner): ______________________  Date: __________',
          'Reviewed by (Farm manager): ______________________  Date: __________',
          'Approved to issue (Farm manager): ______________________  Date: __________',
          'Parameters confirmed (volume / duration / start time): ______________________']:
    d.add_paragraph(l)
d.save('irrigation_order.docx')
