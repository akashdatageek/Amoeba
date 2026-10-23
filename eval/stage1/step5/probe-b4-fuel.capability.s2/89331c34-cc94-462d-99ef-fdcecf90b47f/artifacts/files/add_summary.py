import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
wb=openpyxl.load_workbook('fuel_surcharge_workbook.xlsx')
for n in ('Summary','Sources'):
    if n in wb.sheetnames: del wb[n]
F=Font(name='Arial',size=10); B=Font(name='Arial',size=10,bold=True); BL=Font(name='Arial',size=10,color='0000FF')
inp=PatternFill('solid',fgColor='FFF2CC'); sub=PatternFill('solid',fgColor='F8CBAD'); hd=PatternFill('solid',fgColor='D9E1F2')
s=wb.create_sheet('Summary')
s['A1']='Summary - 52 weeks, EIA Midwest (PADD 2) weekly retail diesel, proxy for Indiana'; s['A1'].font=Font(name='Arial',size=12,bold=True)
for c,h in zip('ABCD',['Metric','Value','Formula / source','As of']): 
    s[c+'3']=h; s[c+'3'].font=B; s[c+'3'].fill=hd
W='Weekly!'
R=lambda c:f"{W}${c}$2:${c}$53"
rows=[
('Average price ($/gal)',f"=AVERAGE({R('B')})",'AVERAGE(Weekly!B2:B53); EIA EMD_EPD2D_PTE_R20_DPG','Weeks ending 2025-10-27 to 2026-10-19','0.000'),
('Average surcharge ($/mi)',f"=AVERAGE({R('C')})",'AVERAGE(Weekly!C2:C53)','same 52 weeks','0.0000'),
('Highest price ($/gal)',f"=MAX({R('B')})",'MAX(Weekly!B2:B53)','same 52 weeks','0.000'),
('Highest week (week ending)',f"=INDEX({R('A')},MATCH(B6,{R('B')},0))",'INDEX/MATCH on highest price (first match if tied)','same 52 weeks','yyyy-mm-dd'),
('Surcharge in highest week ($/mi)',f"=INDEX({R('C')},MATCH(B6,{R('B')},0))",'INDEX/MATCH','same 52 weeks','0.00'),
('Lowest price ($/gal)',f"=MIN({R('B')})",'MIN(Weekly!B2:B53)','same 52 weeks','0.000'),
('Lowest week (week ending)',f"=INDEX({R('A')},MATCH(B9,{R('B')},0))",'INDEX/MATCH on lowest price (first match if tied)','same 52 weeks','yyyy-mm-dd'),
('Surcharge in lowest week ($/mi)',f"=INDEX({R('C')},MATCH(B9,{R('B')},0))",'INDEX/MATCH','same 52 weeks','0.00'),
('52-week total fleet fuel cost ($)',f"=SUM({R('E')})",'SUM(Weekly!E2:E53)','same 52 weeks','#,##0.00'),
('Average weekly fleet fuel cost ($)',f"=AVERAGE({R('E')})",'AVERAGE(Weekly!E2:E53)','same 52 weeks','#,##0.00'),
('Carrier average fuel cost per mile ($/mi)',f"=AVERAGE({R('F')})",'AVERAGE(Weekly!F2:F53); weekly price / 6.2 mpg','same 52 weeks','0.0000'),
('ATRI total marginal cost per mile ($/mi) - INPUT',2.336,'ATRI, An Analysis of the Operational Costs of Trucking (2026 ed.), press release https://truckingresearch.org/2026/07/new-atri-report-details-accelerating-costs-and-low-profitability-despite-cuts [S35][S46]','Cost data year 2025; released 2026-07-15; retrieved 2026-10-23','0.000'),
('ATRI cost per mile excluding fuel ($/mi) - INPUT',1.854,'Same ATRI report and page [S46]','Cost data year 2025; released 2026-07-15; retrieved 2026-10-23','0.000'),
('ATRI fuel cost per mile ($/mi) - SUBSTITUTE',"=B15-B16",'SUBSTITUTE: stated fuel $/mi table value NOT FOUND; derived as total minus non-fuel (2.336 - 1.854 = 0.482)','Cost data year 2025; released 2026-07-15','0.000'),
('ATRI fuel share of total marginal cost',"=B17/B15",'Derived: ATRI fuel $/mi / ATRI total $/mi (uses the substitute above)','Cost data year 2025; released 2026-07-15','0.0%'),
('Gap: carrier minus ATRI fuel per mile ($/mi)',"=B14-B17",'Carrier 52-week average (B14) minus ATRI fuel $/mi (B17)','Carrier: 2025-10-27 to 2026-10-19; ATRI: 2025 data','0.0000'),
('Gap as % of ATRI fuel per mile',"=B19/B17",'B19 / B17','as above','0.0%'),
]
for i,(m,f,src,asof,nf) in enumerate(rows,start=4):
    s.cell(i,1,m).font=F; c=s.cell(i,2,f); c.font=BL if not isinstance(f,str) else F; c.number_format=nf
    s.cell(i,3,src).font=F; s.cell(i,4,asof).font=F
    if not isinstance(f,str): c.fill=inp
    if 'SUBSTITUTE' in m: c.fill=sub; s.cell(i,1).fill=sub
    s.cell(i,3).alignment=Alignment(wrap_text=True,vertical='top')
    s.cell(i,4).alignment=Alignment(wrap_text=True,vertical='top')
s['A22']='Legend: yellow = ATRI input typed in with source and date (blue font); orange = substitute value (derived, not a stated table figure). All other cells are formulas. Carrier $/mi is the 2025-26 price / 6.2 mpg; ATRI is a 2025 industry sample average, so year and basis differ.'
s['A22'].font=Font(name='Arial',size=9,italic=True)
s.column_dimensions['A'].width=46; s.column_dimensions['B'].width=14; s.column_dimensions['C'].width=70; s.column_dimensions['D'].width=40
t=wb.create_sheet('Sources')
t['A1']='Sources'; t['A1'].font=Font(name='Arial',size=12,bold=True)
data=[
('Item','Detail','Retrieved / as of'),
('Proxy statement','EIA publishes no Indiana-only weekly retail diesel series; the EIA weekly Midwest (PADD 2) retail diesel price is used as a proxy for Indiana (an assumption).','Applied 2026-10-23'),
('EIA series ID','EMD_EPD2D_PTE_R20_DPG - Weekly Midwest No. 2 Diesel Retail Prices, dollars per gallon [S1]','Retrieved 2026-10-23'),
('EIA page URL','https://www.eia.gov/dnav/pet/hist/LeafHandler.ashx?n=PET&s=emd_epd2d_pte_r20_dpg&f=W [S1]','Retrieved 2026-10-23'),
('EIA API URL','EIA API v2 petroleum/pri/gnd data route (URL in Weekly!I1) [S64]','Retrieved 2026-10-23'),
('Price period','52 weeks, weeks ending 2025-10-27 to 2026-10-19; latest week as of 2026-10-19 = 6.352 $/gal (FRED shows the same value [S58])','Retrieved 2026-10-23'),
('ATRI report title','An Analysis of the Operational Costs of Trucking, 2026 edition (press headline: New ATRI Report Details Accelerating Costs and Low Profitability Despite Cuts)','Retrieved 2026-10-23'),
('ATRI URL','https://truckingresearch.org/2026/07/new-atri-report-details-accelerating-costs-and-low-profitability-despite-cuts [S35][S46]','Retrieved 2026-10-23'),
('ATRI release date','2026-07-15 (page dateline "Wednesday, July 15, 2026" [S47]; press release says "today released")','as of 2026-07-15'),
('ATRI cost data year','2025; sample 182,248 trucks and 14.67 billion miles [S46]','Data year 2025'),
('ATRI figures used','Total marginal cost 2.336 $/mi (+3.4%); cost excluding fuel 1.854 $/mi (+4.2%) [S46]','as of 2026-07-15 release'),
('ATRI SUBSTITUTE (marked)','Fuel cost per mile stated as a table value NOT FOUND (press release and summary page only; full PDF not read). Substitute: 2.336 - 1.854 = 0.482 $/mi; share 0.482/2.336 = 20.6%.','Derived 2026-10-23'),
('ATRI scope caveat','Industry-wide sample of for-hire/other carriers, all years 2025 data, miles-weighted average marginal cost; not specific to Indiana or a 6.2 mpg Class 8 fleet. The carrier figure uses 2025-26 prices at 6.2 mpg, so year and basis differ.','Caveat'),
('ATRI newer edition','No newer edition than July 2026 was seen in the latest search (results showed only the 2026 edition and an older report article); the newest edition found is used. Only search snippets were seen for this check.','Checked 2026-10-23'),
('Carrier parameters','Base 1.25 $/gal, 6.2 mpg, 35 tractors, 2,300 miles per tractor per week: given in the task (Inputs sheet)','Task'),
('Surcharge method','ROUND((price - 1.25) / 6.2, 2), every loaded and empty mile','Task'),
]
for r,row in enumerate(data,start=3):
    for c,v in enumerate(row,start=1):
        x=t.cell(r,c,v); x.font=B if r==3 else F; x.alignment=Alignment(wrap_text=True,vertical='top')
        if r==3: x.fill=hd
        if 'SUBSTITUTE' in str(row[0]): x.fill=sub
t.column_dimensions['A'].width=26; t.column_dimensions['B'].width=110; t.column_dimensions['C'].width=26
wb.save('fuel_surcharge_workbook.xlsx')
