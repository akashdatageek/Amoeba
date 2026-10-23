import csv
from decimal import Decimal as D, ROUND_HALF_UP
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
rows=[(r['week_ending'],D(r['price_usd_per_gal'])) for r in csv.DictReader(open('weekly_prices.csv'))]
assert len(rows)==52
gal=D(35*2300)/D('6.2')
q=lambda x,p='0.01':x.quantize(D(p),ROUND_HALF_UP)
sur=lambda p:q((p-D('1.25'))/D('6.2'))
tot=sum(p*gal for _,p in rows); avgp=sum(p for _,p in rows)/52
hi=max(rows,key=lambda r:r[1]); lo=min(rows,key=lambda r:r[1])
avgs=sum(sur(p) for _,p in rows)/52; fpm=avgp/D('6.2'); gap=fpm-D('0.482')
print(q(tot),q(avgp,'0.00001'),hi,lo,q(avgs,'0.0001'),q(fpm,'0.0001'),q(gap,'0.0001'),q(gap/D('0.482')*100,'0.1'))
d=Document(); s=d.sections[0]
s.page_width=Inches(8.5); s.page_height=Inches(11)
for a in('left_margin','right_margin','top_margin','bottom_margin'): setattr(s,a,Inches(0.7))
st=d.styles['Normal']; st.font.name='Calibri'; st.font.size=Pt(10)
st.element.rPr.rFonts.set(qn('w:eastAsia'),'Calibri')
st.paragraph_format.space_after=Pt(2); st.paragraph_format.space_before=Pt(0)
def P(t,b=False,sz=None,al=None):
    p=d.add_paragraph(); r=p.add_run(t); r.bold=b
    if sz: r.font.size=Pt(sz)
    return p
def H(t):
    p=P(t,True,11); p.paragraph_format.space_before=Pt(5); return p
def B(t):
    p=d.add_paragraph(style='List Bullet'); p.add_run(t); p.paragraph_format.space_after=Pt(0); return p
def shade(c):
    tcPr=c._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:fill'),'D9E2F3'); tcPr.append(sh)
def table(data,widths,hdr=True,sz=9):
    t=d.add_table(rows=len(data),cols=len(data[0])); t.style='Table Grid'; t.autofit=False
    for i,row in enumerate(data):
        for j,v in enumerate(row):
            c=t.cell(i,j); c.width=Inches(widths[j]); p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
            r=p.add_run(str(v)); r.font.size=Pt(sz)
            if hdr and i==0: r.bold=True; shade(c)
    return t
P('I read "weekly retail on-highway diesel price for Indiana" as EIA weekly Midwest (PADD 2) retail diesel price, used as a proxy because EIA publishes no Indiana-only weekly series; if you meant EIA weekly Indiana-specific retail diesel price, if such a series exists or EIA weekly Midwest regional price, with Indianapolis or a nearby metro as a cross-check or EIA weekly US national average diesel price, as the fallback, the answer below may not apply.',False,8.5).runs[0].italic=True
P('MEMORANDUM',True,13)
P('To: Management, Indianapolis freight carrier     From: Delivery lead, fuel analysis     Date: 23 October 2026\nSubject: Fuel-surcharge schedule and 52-week fuel cost (35 Class 8 tractors)')
H('1. Bottom line')
P(f'Using the standard formula, the surcharge over the last 52 weeks averaged ${q(avgs,"0.00")} per mile on an average Midwest diesel price of ${q(avgp,"0.00")} per gallon [S3]. Fuel cost the fleet ${q(tot,"0.01"):,} over the 52 weeks, or ${q(fpm,"0.00")} per mile, which is ${q(gap,"0.00")} per mile above the ATRI fuel figure of $0.48 per mile derived below [S35]. Prices ran above the $5.00 top of the schedule grid in recent weeks (latest week $6.352), so the grid needs extending upward.')
H('2. Data and method')
B('Price series: EIA Midwest (PADD 2) No. 2 Diesel Retail Price, weekly, $/gal, series EMD_EPD2D_PTE_R20_DPG; EIA API v2, retrieved 2026-10-23 [S3]. 52 weeks, Monday price dates 2025-10-27 to 2026-10-19 (latest week = $6.352 as of 2026-10-19). The EIA dnav page copy was released 10/14/2026 and ends at 10/12 [S1]. Midwest stands in for Indiana (see section 6).')
B('Surcharge per mile = (price − $1.25) ÷ 6.2 mpg, rounded to the cent (half-up), applied to every loaded and empty mile. Fleet: 35 tractors × 2,300 mi ÷ 6.2 mpg = 12,983.87 gal per week; weekly fleet fuel cost = price × 12,983.87.')
B('Assumptions: constant weekly gallons (no change in miles or mpg); each 5-cent band is priced at its midpoint; surcharge is on every mile. Weekly detail and formulas are in the attached workbook fuel_surcharge_workbook.xlsx; the price and surcharge chart is the attached price_surcharge_chart.png.')
H('3. Surcharge schedule, $3.00 to $5.00 in 5-cent bands (cents per mile = $)')
P('Each band uses ROUND((midpoint − 1.25) ÷ 6.2, 2); e.g. $3.00–3.05 uses 3.025. Computed from the formula and task parameters; the base and mpg are from the task. Above $5.00 the same formula applies (e.g. $6.68 gives $0.88).',False,9)
bands=[]
for i in range(40):
    lo_=D('3.00')+D('0.05')*i; mid=lo_+D('0.025')
    bands.append((f'${lo_:.2f}–{lo_+D("0.05"):.2f}',f'${sur(mid):.2f}'))
data=[['Price band ($/gal)','$/mile']*4]
for r in range(10):
    data.append([x for c in range(4) for x in bands[c*10+r]])
table(data,[1.0,0.75]*4)
H('4. 52-week results (EIA series above, retrieved 2026-10-23 [S3])')
fc=lambda p:f'${q(p*gal,"0.01"):,}'
table([['Measure','Week (Monday date)','Price $/gal','Surcharge $/mi','Fleet fuel cost'],
 ['Average (52 weeks)','—',f'${q(avgp,"0.000")}',f'${q(avgs,"0.00")}',f'${q(tot/52,"0.01"):,} per week'],
 ['Highest',hi[0],f'${hi[1]}',f'${sur(hi[1])}',fc(hi[1])],
 ['Lowest',lo[0],f'${lo[1]}',f'${sur(lo[1])}',fc(lo[1])]],[1.5,1.5,1.0,1.2,1.9])
P(f'Total 52-week fleet fuel cost: ${q(tot,"0.01"):,}. Average fuel cost per mile = average of weekly (price ÷ 6.2) = ${q(fpm,"0.0000")}. The average surcharge is the mean of the rounded weekly values.',False)
H('5. Comparison with ATRI')
P('ATRI, "An Analysis of the Operational Costs of Trucking," press release dated 15 Jul 2026, cost data year 2025: total marginal cost $2.336 per mile and $1.854 per mile excluding fuel [S35]. Fuel is derived as 2.336 − 1.854 = $0.482 per mile, a share of 20.6% (0.482 ÷ 2.336); this is derived, not a published table value. The carrier averaged $'+f'{q(fpm,"0.0000")} per mile, ${q(gap,"0.0000")} above ATRI, or {q(gap/D("0.482")*100,"0.0")}% higher. Caveat: not like-for-like. ATRI is an industry-wide average on 2025 data; the carrier figure uses Midwest retail prices for 2025-10-27 to 2026-10-19 at 6.2 mpg, and the price level rose during 2026 (high $6.68 in the week of 2026-09-21).')
H('6. Not obtained / substituted')
B('Indiana-only weekly diesel price: EIA publishes none; the Midwest (PADD 2) series is the proxy (task assumption).')
B('ATRI stated fuel cost per mile: not read from the full report PDF (only the press release and summary were read); substituted by total minus non-fuel cost = $0.482. No search for an ATRI edition newer than July 2026 was run, so that point is unconfirmed.')
B('Live re-fetch of EIA and ATRI for verification was not possible (no network); checks used pages saved 2026-10-23.')
H('7. Sources')
for t in ['[S3] EIA API v2, petroleum/pri/gnd, series EMD_EPD2D_PTE_R20_DPG, weekly, retrieved 2026-10-23: https://api.eia.gov/v2/petroleum/pri/gnd/data/',
 '[S1] EIA, Weekly Midwest No 2 Diesel Retail Prices, released 10/14/2026: https://www.eia.gov/dnav/pet/hist/LeafHandler.ashx?n=PET&s=emd_epd2d_pte_r20_dpg&f=W',
 '[S35] ATRI, "New ATRI Report Details Accelerating Costs and Low Profitability Despite Cuts," 15 Jul 2026: https://truckingresearch.org/2026/07/new-atri-report-details-accelerating-costs-and-low-profitability-despite-cuts',
 '[S58] FRED GASDESMWW (EIA Midwest diesel, cross-check of 2026-10-19), retrieved 2026-10-23: https://fred.stlouisfed.org/series/GASDESMWW']:
    B(t).runs[0].font.size=Pt(9)
d.save('fuel_surcharge_memo.docx')
