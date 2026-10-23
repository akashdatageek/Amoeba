import csv, datetime as dt
from decimal import Decimal, ROUND_HALF_UP
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment
rows=list(csv.DictReader(open('weekly_prices.csv',newline='')))
assert len(rows)==52
F=lambda **k: Font(**{'name':'Arial','size':10,**k})
H=PatternFill('solid',fgColor='1F3864'); INP=PatternFill('solid',fgColor='FFF2CC')
wb=Workbook()
ws=wb.active; ws.title='Inputs'
ws['A1']='Fleet fuel surcharge model - Inputs'; ws['A1'].font=F(bold=True,size=13)
ws['A2']='Yellow cells are the only hard-coded parameters; edit them and every sheet recalculates.'; ws['A2'].font=F(italic=True)
P=[('Base price ($/gal)',1.25,'Standard base price in the task'),('Fleet fuel economy (mpg)',6.2,'Given in the task: Class 8 tractors average 6.2 mpg'),('Tractors (count)',35,'Given in the task'),('Miles per tractor per week',2300,'Given in the task')]
ws['A3'],ws['B3'],ws['C3']='Parameter','Value','Basis'
for c in 'ABC': ws[c+'3'].font=F(bold=True,color='FFFFFF'); ws[c+'3'].fill=H
for i,(a,b,c) in enumerate(P):
    r=4+i; ws.cell(r,1,a); ws.cell(r,2,b).fill=INP; ws.cell(r,3,c)
ws['A8']='Fleet miles per week (derived)'; ws['B8']='=B6*B7'; ws['C8']='Tractors x miles per tractor'
ws['A9']='Fleet gallons per week (derived)'; ws['B9']='=B8/B5'; ws['C9']='Fleet miles divided by mpg'
ws['B8'].number_format='#,##0'; ws['B9'].number_format='#,##0.0'
ws['A11']='Assumption notes'; ws['A11'].font=F(bold=True)
notes=['1. Constant gallons: every week uses the same gallons (35 x 2,300 / 6.2); miles and mpg do not vary by week.',
'2. Midpoint band convention: each 5-cent band surcharge is computed at the band midpoint, ROUND((midpoint - base)/mpg, 2).',
'3. Surcharge applies to every mile, loaded and empty: Surcharge per mile = ROUND((price - base) / mpg, 2).',
'4. Class 8 = heavy-duty truck with GVWR over 33,000 lb (semi tractor).',
'5. Price proxy: EIA weekly Midwest (PADD 2) No. 2 diesel retail price (series EMD_EPD2D_PTE_R20_DPG); EIA publishes no Indiana-only weekly series.',
'6. Fleet fuel cost = price x weekly gallons; fuel cost per mile = price / mpg (no tax or discount adjustments).',
'7. Band grid (start $3.00, 5-cent step) is on the Bands sheet as labelled grid settings from the task, not a carrier parameter.']
for i,n in enumerate(notes): ws.cell(12+i,1,n)
ws.column_dimensions['A'].width=40; ws.column_dimensions['B'].width=14; ws.column_dimensions['C'].width=60
for row in ws.iter_rows():
    for c in row:
        if c.font.name!='Arial': c.font=F()
for r in range(4,8): ws.cell(r,1).font=F()
# Weekly
w=wb.create_sheet('Weekly')
hd=['Week ending','Price $/gal','Surcharge $/mi','Gallons','Fleet fuel cost $','Fuel cost $/mi','Source']
for j,h in enumerate(hd,1):
    c=w.cell(1,j,h); c.font=F(bold=True,color='FFFFFF'); c.fill=H; c.alignment=Alignment(wrap_text=True,horizontal='center')
for i,r in enumerate(rows):
    n=i+2
    w.cell(n,1,dt.date.fromisoformat(r['week_ending'])).number_format='yyyy-mm-dd'
    w.cell(n,2,float(r['price_usd_per_gal'])).number_format='0.000'
    w.cell(n,3,f'=ROUND((B{n}-Inputs!$B$4)/Inputs!$B$5,2)').number_format='0.00'
    w.cell(n,4,'=Inputs!$B$9').number_format='#,##0.0'
    w.cell(n,5,f'=B{n}*D{n}').number_format='#,##0.00'
    w.cell(n,6,f'=B{n}/Inputs!$B$5').number_format='0.0000'
    w.cell(n,7,f"EIA series {r['series_id']} (Midwest PADD 2 weekly retail diesel), retrieved {r['retrieved_date']}")
    for j in range(1,8): w.cell(n,j).font=F()
w['I1']='Prices (column B) are values from weekly_prices.csv; all other columns are formulas. API URL: '+rows[0]['source_url']
w['I1'].font=F(italic=True)
for col,wd in zip('ABCDEFG',[13,12,14,11,17,14,70]): w.column_dimensions[col].width=wd
w.freeze_panes='A2'
# Bands
b=wb.create_sheet('Bands')
for j,h in enumerate(['Band','Low $/gal','High $/gal','Midpoint $/gal','Surcharge $/mi'],1):
    c=b.cell(1,j,h); c.font=F(bold=True,color='FFFFFF'); c.fill=H
b['G1']='Band grid start ($/gal)'; b['H1']=3.00; b['G2']='Band width ($/gal)'; b['H2']=0.05
b['G3']='Grid settings from the task ($3.00 to $5.00, 5-cent bands); surcharge at midpoint.'
b['H1'].fill=INP; b['H2'].fill=INP
for i in range(40):
    n=i+2
    b.cell(n,1,i+1)
    b.cell(n,2,'=$H$1' if i==0 else f'=C{n-1}')
    b.cell(n,3,f'=ROUND(B{n}+$H$2,2)')
    b.cell(n,4,f'=ROUND((B{n}+C{n})/2,3)')
    b.cell(n,5,f'=ROUND((D{n}-Inputs!$B$4)/Inputs!$B$5,2)')
    for j in range(2,5): b.cell(n,j).number_format='0.000' if j==4 else '0.00'
    b.cell(n,5).number_format='0.00'
    for j in range(1,6): b.cell(n,j).font=F()
for c in ['G1','H1','G2','H2','G3']: b[c].font=F()
for col,wd in zip('ABCDEGH',[8,12,12,16,16,26,10]): b.column_dimensions[col].width=wd
wb.save('fuel_surcharge_workbook.xlsx')
