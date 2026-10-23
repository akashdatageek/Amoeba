import openpyxl,csv,re,struct
from decimal import Decimal as D, ROUND_HALF_UP
def r2(x,n=2): return x.quantize(D(1).scaleb(-n),rounding=ROUND_HALF_UP)
wb=openpyxl.load_workbook('fuel_surcharge_workbook.xlsx')
v=openpyxl.load_workbook('fuel_surcharge_workbook.xlsx',data_only=True)
print(wb.sheetnames)
for n in wb.sheetnames:
    ws=wb[n]; nf=sum(1 for r in ws.iter_rows() for c in r if isinstance(c.value,str) and c.value.startswith('='))
    print(n,ws.max_row,ws.max_column,'formulas',nf)
ws=wb['Inputs']
for r in ws.iter_rows(min_row=1,max_row=20):
    print([c.value for c in r][:3])
W=wb['Weekly'];WV=v['Weekly']
print([c.value for c in W[1]][:7]); print([c.value for c in W[2]][:7]); print([c.value for c in W[53]][:7])
# hardcoded numbers in Weekly C-F, Bands E, Summary
for n,cols in (('Weekly','CDEF'),('Bands','E'),('Summary','B')):
    for row in wb[n].iter_rows(min_row=2):
        for c in row:
            if c.column_letter in cols and c.value is not None and not (isinstance(c.value,str) and c.value.startswith('=')) and isinstance(c.value,(int,float)):
                print('HARDCODE',n,c.coordinate,c.value)
# recompute weekly
csvr=list(csv.DictReader(open('weekly_prices.csv')))
base,mpg=D('1.25'),D('6.2'); gal=D(35*2300)/mpg
mism=0; prices=[];sur=[];cost=[];fpm=[]
for i,row in enumerate(csvr):
    r=i+2; p=D(row['price_usd_per_gal'])
    s=r2((p-base)/mpg); g=gal; fc=p*g; f=p/mpg
    prices.append(p);sur.append(s);cost.append(fc);fpm.append(f)
    wp=D(str(WV.cell(r,2).value)); 
    got=(WV.cell(r,3).value,WV.cell(r,4).value,WV.cell(r,5).value,WV.cell(r,6).value)
    d=WV.cell(r,1).value
    ok = wp==p and str(d)[:10]==row['week_ending'] and abs(D(str(got[0]))-s)<D('0.000001') and abs(D(str(got[1]))-g)<D('0.001') and abs(D(str(got[2]))-fc)<D('0.01') and abs(D(str(got[3]))-f)<D('0.00001')
    fs=[W.cell(r,c).value for c in (3,4,5,6)]
    exp=[f'=ROUND((B{r}-Inputs!$B$4)/Inputs!$B$5,2)']
    if not all(isinstance(x,str) and x.startswith('=') for x in fs): print('nonformula row',r,fs)
    if not ok: mism+=1; print('MISMATCH',r,row['week_ending'],p,s,got)
print('weekly mismatches',mism,'rows',len(csvr))
# spacing
import datetime as dt
ds=[dt.date.fromisoformat(x['week_ending']) for x in csvr]; print('gaps',{(b-a).days for a,b in zip(ds,ds[1:])},ds[0],ds[-1])
# bands
B=wb['Bands'];BV=v['Bands'];bm=0
for k in range(40):
    r=k+2; lo=D('3.00')+D('0.05')*k; hi=lo+D('0.05'); mid=(lo+hi)/2; s=r2((mid-base)/mpg)
    got=[BV.cell(r,c).value for c in range(1,6)]
    if not (abs(D(str(got[1]))-lo)<D('1e-9') and abs(D(str(got[2]))-hi)<D('1e-9') and abs(D(str(got[3]))-mid)<D('1e-9') and abs(D(str(got[4]))-s)<D('1e-9')):
        bm+=1; print('BAND MISMATCH',r,got,lo,hi,mid,s)
print('band mismatches',bm, 'band row2 formulas',[B.cell(2,c).value for c in range(1,6)],'row41',[B.cell(41,c).value for c in range(1,6)])
# summary
S=wb['Summary'];SV=v['Summary']
for r in range(1,24):
    print(r,SV.cell(r,1).value,'|',S.cell(r,2).value,'|',SV.cell(r,2).value,'|',str(SV.cell(r,3).value)[:70],'|',str(SV.cell(r,4).value)[:50])
n=len(prices)
avgp=sum(prices)/n; avgs=sum(sur)/n; tot=sum(cost); avgf=sum(fpm)/n
hi=max(prices);lo=min(prices)
print('EXP avgp',avgp,'avgs',avgs,'tot',r2(tot),'avgweek',r2(tot/n),'avgfpm',avgf,'hi',hi,ds[prices.index(hi)],sur[prices.index(hi)],'lo',lo,ds[prices.index(lo)],sur[prices.index(lo)])
atri=D('2.336');nf=D('1.854');fu=atri-nf
print('EXP atri fuel',fu,'share',fu/atri,'gap',avgf-fu,'gap%',(avgf-fu)/fu)
# Sources
So=wb['Sources']
for r in So.iter_rows():
    vals=[str(c.value)[:110] for c in r if c.value is not None]
    if vals: print(r[0].row,vals)
# png
d=open('price_surcharge_chart.png','rb').read(); print('PNG',struct.unpack('>II',d[16:24]))
