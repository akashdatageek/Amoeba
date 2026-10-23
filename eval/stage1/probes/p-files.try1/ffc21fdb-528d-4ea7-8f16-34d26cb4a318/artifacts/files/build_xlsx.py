from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
wb = Workbook()
ws = wb.active
ws.title = 'Prices'
ws.append(['Item', 'Qty', 'Unit price', 'Total'])
ws.append(['Pens', 10, 1.5, '=B2*C2'])
ws.append(['Pads', 4, 3.25, '=B3*C3'])
ws.append(['Clips', 100, 0.02, '=B4*C4'])
ws['A5'] = 'Grand total'
ws['D5'] = '=SUM(D2:D4)'
bold = Font(bold=True)
for c in ws[1]: c.font = bold
for c in ws[5]: c.font = bold
for r in range(2, 6):
    ws.cell(r, 3).number_format = '0.00'
    ws.cell(r, 4).number_format = '0.00'
for col, w in zip('ABCD', (16, 10, 12, 12)):
    ws.column_dimensions[col].width = w
wb.save('prices.xlsx')
