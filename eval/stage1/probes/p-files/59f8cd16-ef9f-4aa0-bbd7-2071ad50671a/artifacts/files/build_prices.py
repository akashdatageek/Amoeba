from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
wb = Workbook()
ws = wb.active
ws.title = "Prices"
ws.append(["Item", "Qty", "Unit price", "Total"])
rows = [("pens", 10, 1.50), ("pads", 4, 3.25), ("clips", 100, 0.02)]
for i, (n, q, p) in enumerate(rows, start=2):
    ws.append([n, q, p, f"=B{i}*C{i}"])
ws["A5"] = "Grand total"
ws["D5"] = "=SUM(D2:D4)"
for r in range(2, 6):
    ws[f"C{r}"].number_format = "0.00"
    ws[f"D{r}"].number_format = "0.00"
for c in ws[1]:
    c.font = Font(bold=True)
for c in ws[5]:
    c.font = Font(bold=True)
for col, w in zip("ABCD", (16, 8, 12, 12)):
    ws.column_dimensions[col].width = w
wb.save("prices.xlsx")
wb2 = load_workbook("prices.xlsx")
s = wb2["Prices"]
print(wb2.sheetnames)
for row in s.iter_rows():
    print([c.coordinate + "=" + repr(c.value) for c in row])
print("formulas ok:", all(isinstance(s[f"D{r}"].value, str) and s[f"D{r}"].value.startswith("=") for r in range(2, 6)))
