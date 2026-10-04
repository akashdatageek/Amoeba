import openpyxl
from openpyxl.styles import Font

# Create a new workbook and select the active sheet
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Assignment Plan"

# --- 1. Loads Table ---
ws["A1"] = "Load ID"
ws["B1"] = "Weight (lb)"
ws["C1"] = "Miles"
for cell in ws[1]:
    cell.font = Font(bold=True)

loads_data = [
    ("L1", 18000, 220),
    ("L2", 12500, 340),
    ("L3", 22000, 150),
    ("L4", 9000, 410),
    ("L5", 15500, 95),
    ("L6", 7000, 280),
    ("L7", 20500, 190),
    ("L8", 11000, 365),
    ("L9", 14000, 120),
    ("L10", 8500, 255),
]

for i, (lid, weight, miles) in enumerate(loads_data, start=2):
    ws.cell(row=i, column=1, value=lid)
    ws.cell(row=i, column=2, value=weight)
    ws.cell(row=i, column=3, value=miles)

# --- 2. Trucks Table ---
ws["A13"] = "Truck ID"
ws["B13"] = "Capacity (lb)"
ws["C13"] = "Mileage Limit"
ws["D13"] = "Fixed Cost ($)"
ws["E13"] = "Rate ($/mile)"
for cell in ws[13]:
    cell.font = Font(bold=True)

trucks_data = [
    ("A", 45000, 1100, 350, 2.10),
    ("B", 45000, 1100, 350, 2.10),
    ("C", 30000, 1100, 500, 1.85),
    ("D", 26000, 900, 275, 2.45),
]

for i, (tid, cap, lim, fix, rate) in enumerate(trucks_data, start=14):
    ws.cell(row=i, column=1, value=tid)
    ws.cell(row=i, column=2, value=cap)
    ws.cell(row=i, column=3, value=lim)
    ws.cell(row=i, column=4, value=fix)
    ws.cell(row=i, column=5, value=rate)

# --- 3. Assignment Table ---
ws["A19"] = "Truck ID"
ws["B19"] = "Loads"
ws["C19"] = "Total Weight (lb)"
ws["D19"] = "Total Miles"
for cell in ws[19]:
    cell.font = Font(bold=True)

assignment_data = [
    ("A", "L1, L2, L9", 44500, 680),
    ("B", "L5, L7, L10", 44500, 540),
    ("C", "L4, L6, L8", 27000, 1055),
    ("D", "L3", 22000, 150),
]

for i, (tid, loads, weight, miles) in enumerate(assignment_data, start=20):
    ws.cell(row=i, column=1, value=tid)
    ws.cell(row=i, column=2, value=loads)
    ws.cell(row=i, column=3, value=weight)
    ws.cell(row=i, column=4, value=miles)

# --- 4. Cost Calculation Table ---
ws["A25"] = "Truck ID"
ws["B25"] = "Fixed Cost ($)"
ws["C25"] = "Rate ($/mile)"
ws["D25"] = "Total Miles"
ws["E25"] = "Truck Cost ($)"
for cell in ws[25]:
    cell.font = Font(bold=True)

# Formulas for each truck
# Truck A: Fixed=D14, Rate=E14, Miles=D20
# Truck B: Fixed=D15, Rate=E15, Miles=D21
# Truck C: Fixed=D16, Rate=E16, Miles=D22
# Truck D: Fixed=D17, Rate=E17, Miles=D23

cost_rows = [
    ("A", "D14", "E14", "D20"),
    ("B", "D15", "E15", "D21"),
    ("C", "D16", "E16", "D22"),
    ("D", "D17", "E17", "D23"),
]

for i, (tid, fix_ref, rate_ref, miles_ref) in enumerate(cost_rows, start=26):
    ws.cell(row=i, column=1, value=tid)
    ws.cell(row=i, column=2, value=f"={fix_ref}")
    ws.cell(row=i, column=3, value=f"={rate_ref}")
    ws.cell(row=i, column=4, value=f"={miles_ref}")
    # Formula: Fixed + (Rate * Miles)
    ws.cell(row=i, column=5, value=f"=B{i}+(C{i}*D{i})")

ws["A30"] = "Total Cost"
ws["E30"] = "=SUM(E26:E29)"
ws["A30"].font = Font(bold=True)
ws["E30"].font = Font(bold=True)

# Save the workbook
wb.save("assignment_plan.xlsx")
