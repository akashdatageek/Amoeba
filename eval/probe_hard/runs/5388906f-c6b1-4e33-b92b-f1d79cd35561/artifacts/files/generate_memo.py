from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_memo():
    doc = Document()

    # Set style for the whole document
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)

    # Memo Header
    header = doc.add_paragraph()
    header.add_run('MEMORANDUM').bold = True
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.runs[0].font.size = Pt(14)

    memo_info = doc.add_paragraph()
    memo_info.add_run("TO: Management\n").bold = True
    memo_info.add_run("FROM: Document Specialist\n").bold = True
    memo_info.add_run("DATE: October 4, 2026\n").bold = True
    memo_info.add_run("SUBJECT: Total Cost of Ownership (TCO) Analysis: Diesel vs. Battery-Electric Class 8 Tractors").bold = True
    
    doc.add_paragraph("-" * 50)

    # Introduction
    doc.add_paragraph("This memo provides a 7-year Total Cost of Ownership (TCO) comparison between traditional diesel Class 8 tractors and battery-electric (BEV) alternatives for our regional routes (under 250 miles/day, 60,000 miles/year per truck).")

    # TCO Table
    doc.add_paragraph("\n7-Year TCO Comparison (Per Truck)").bold = True
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Cost Component'
    hdr_cells[1].text = 'Diesel Tractor'
    hdr_cells[2].text = 'Battery-Electric (BEV)'

    data = [
        ("Purchase Price", "$185,000 [S2]", "$420,000 [S2]"),
        ("Federal Incentives", "$0 [S2]", "$0 [S2]"),
        ("Indiana Incentives", "$0 [unverified]", "$0 [unverified]"),
        ("Energy Cost (420k miles)", "$228,000", "$100,800"),
        ("Maintenance (420k miles)", "$63,000", "$37,800"),
        ("Charging Infrastructure", "$0", "$50,000 [unverified]"),
        ("Total 7-Year TCO", "$476,000", "$608,600"),
    ]

    for item, diesel, bev in data:
        row_cells = table.add_row().cells
        row_cells[0].text = item
        row_cells[1].text = diesel
        row_cells[2].text = bev
        if item == "Total 7-Year TCO":
            row_cells[0].paragraphs[0].runs[0].bold = True
            row_cells[1].paragraphs[0].runs[0].bold = True
            row_cells[2].paragraphs[0].runs[0].bold = True

    # Assumptions and Citations
    doc.add_paragraph("\nAssumptions and Citations").bold = True
    assumptions = [
        "Mileage: 60,000 miles/year for 7 years (Total: 420,000 miles).",
        "Diesel Efficiency: 7.0 MPG [S1].",
        "Diesel Fuel Price: $3.80/gallon [unverified].",
        "BEV Energy Efficiency: 2.0 kWh/mi [S1].",
        "Commercial Electricity Rate: $0.12/kWh [unverified].",
        "Diesel Maintenance Cost: $0.15/mile [unverified].",
        "BEV Maintenance Cost: $0.09/mile [unverified].",
        "Federal Incentives: Confirmed as $0 due to cancellation of IRA 45W tax credit [S2].",
        "Infrastructure: Estimated $50,000 per truck for charging installation [unverified]."
    ]
    for assumption in assumptions:
        doc.add_paragraph(assumption, style='List Bullet')

    # Recommendation
    doc.add_paragraph("\nRecommendation").bold = True
    rec_text = (
        "Based on the current financial analysis, it is recommended to continue using diesel tractors. "
        "The Battery-Electric alternative results in a significantly higher Total Cost of Ownership, "
        "costing an additional $132,600 per truck over a 7-year period. The higher purchase price "
        "and infrastructure costs are not sufficiently offset by energy and maintenance savings, "
        "especially given the lack of available federal and state incentives in 2026."
    )
    doc.add_paragraph(rec_text)

    doc.save('TCO_Recommendation_Memo.docx')

if __name__ == "__main__":
    create_memo()
