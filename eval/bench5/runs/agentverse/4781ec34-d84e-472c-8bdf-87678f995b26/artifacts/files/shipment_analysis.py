import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side

# Create workbook and sheet
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Shipment Data"

# Define headers
headers = ["Lane", "Loads", "Avg Revenue per Load", "Total Revenue", "Volume Contribution %", "Revenue Contribution %"]
ws.append(headers)

# Define data
data = [
    ["Lane A", 12, 1850],
    ["Lane B", 7, 2400],
    ["Lane C", 20, 1420],
    ["Lane D", 4, 3100],
]

# Populate data and add formulas for Total Revenue
for row_idx, row_data in enumerate(data, start=2):
    ws.cell(row=row_idx, column=1, value=row_data[0])
    ws.cell(row=row_idx, column=2, value=row_data[1])
    ws.cell(row=row_idx, column=3, value=row_data[2])
    # Total Revenue = Loads * Avg Revenue
    ws.cell(row=row_idx, column=4, value=f"=B{row_idx}*C{row_idx}")
    # Volume Contribution % = Loads / Total Loads
    ws.cell(row=row_idx, column=5, value=f"=B{row_idx}/$B$6")
    # Revenue Contribution % = Total Revenue / Total Revenue
    ws.cell(row=row_idx, column=6, value=f"=D{row_idx}/$D$6")

# Totals Row (Row 6)
total_row = 6
ws.cell(row=total_row, column=1, value="Total")
ws.cell(row=total_row, column=2, value=f"=SUM(B2:B5)")
ws.cell(row=total_row, column=3, value=f"=IFERROR(D{total_row}/B{total_row}, 0)")
ws.cell(row=total_row, column=4, value=f"=SUM(D2:D5)")

# Formatting
font_arial = Font(name="Arial")
font_arial_bold = Font(name="Arial", bold=True)

# Borders
thin = Side(border_style="thin", color="000000")
border = Border(top=thin, left=thin, right=thin, bottom=thin)

# Alignment
align_left = Alignment(horizontal="left")
align_center = Alignment(horizontal="center")
align_right = Alignment(horizontal="right")

# Apply formatting to all cells in the table
for row in ws.iter_rows(min_row=1, max_row=total_row, min_col=1, max_col=6):
    for cell in row:
        cell.font = font_arial
        cell.border = border
        
        # Column-based alignment
        if cell.column == 1:
            cell.alignment = align_left
        elif cell.column == 2:
            cell.alignment = align_center
        else:
            cell.alignment = align_right

# Header and Total row bolding
for cell in ws[1]:
    cell.font = font_arial_bold
for cell in ws[total_row]:
    cell.font = font_arial_bold

# Number Formatting
# Currency: Avg Revenue and Total Revenue
for row in range(2, total_row + 1):
    ws.cell(row=row, column=3).number_format = '$#,##0.00'
    ws.cell(row=row, column=4).number_format = '$#,##0.00'

# Percentage: Contribution columns
for row in range(2, total_row):
    ws.cell(row=row, column=5).number_format = '0.00%'
    ws.cell(row=row, column=6).number_format = '0.00%'

# Adjust column widths for readability
ws.column_dimensions['A'].width = 15
ws.column_dimensions['B'].width = 10
ws.column_dimensions['C'].width = 22
ws.column_dimensions['D'].width = 22
ws.column_dimensions['E'].width = 22
ws.column_dimensions['F'].width = 22

wb.save("shipment_data.xlsx")
