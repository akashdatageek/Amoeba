import openpyxl
from openpyxl.styles import Font, Border, Side, Alignment

# Create workbook and sheet
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Shipment Data"

# Data
data = [
    ["Lane A", 12, 1850],
    ["Lane B", 7, 2400],
    ["Lane C", 20, 1420],
    ["Lane D", 4, 3100],
]

# Headers
headers = ["Lane", "Loads", "Avg Revenue per Load", "Total Revenue", "Revenue Contribution %"]
ws.append(headers)

# Fill data and add formulas for Total Revenue
for row_idx, row_data in enumerate(data, start=2):
    ws.cell(row=row_idx, column=1, value=row_data[0])
    ws.cell(row=row_idx, column=2, value=row_data[1])
    ws.cell(row=row_idx, column=3, value=row_data[2])
    # Total Revenue = Loads * Avg Revenue
    ws.cell(row=row_idx, column=4, value=f"=B{row_idx}*C{row_idx}")
    # Contribution % = Total Revenue / Total Revenue Sum (D6)
    ws.cell(row=row_idx, column=5, value=f"=D{row_idx}/$D$6")

# Totals Row (Row 6)
total_row = 6
ws.cell(row=total_row, column=1, value="Total")
ws.cell(row=total_row, column=2, value=f"=SUM(B2:B5)")
ws.cell(row=total_row, column=3, value=f"=IFERROR(D{total_row}/B{total_row}, 0)")
ws.cell(row=total_row, column=4, value=f"=SUM(D2:D5)")
ws.cell(row=total_row, column=5, value="")

# Formatting
font_arial = Font(name="Arial")
font_bold_arial = Font(name="Arial", bold=True)
side = Side(style="thin")
border = Border(left=side, right=side, top=side, bottom=side)

for row in ws.iter_rows(min_row=1, max_row=total_row, min_col=1, max_col=5):
    for cell in row:
        cell.font = font_arial
        cell.border = border
        cell.alignment = Alignment(horizontal="center")

for col in range(1, 6):
    ws.cell(row=1, column=col).font = font_bold_arial
    ws.cell(row=total_row, column=col).font = font_bold_arial

for row in range(2, total_row + 1):
    ws.cell(row=row, column=3).number_format = "\$#,##0.00"
    ws.cell(row=row, column=4).number_format = "\$#,##0.00"

for row in range(2, total_row):
    ws.cell(row=row, column=5).number_format = "0.00%"

ws.column_dimensions['A'].width = 15
ws.column_dimensions['B'].width = 10
ws.column_dimensions['C'].width = 22
ws.column_dimensions['D'].width = 20
ws.column_dimensions['E'].width = 25

wb.save("shipment_data.xlsx")