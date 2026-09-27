import openpyxl
from openpyxl.styles import Font

# Create a new workbook and select the active sheet
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Shipments"

# Define the professional font
professional_font = Font(name='Arial')

# 1. Table Structure - Headers
headers = ['Lane', 'Loads', 'Avg Revenue', 'Total Revenue']
for col_num, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_num, value=header)
    cell.font = professional_font

# 2. Row Data
data = [
    ('Lane A', 12, 1850),
    ('Lane B', 7, 2400),
    ('Lane C', 20, 1420),
    ('Lane D', 4, 3100),
]

for row_idx, (lane, loads, avg_rev) in enumerate(data, 2):
    ws.cell(row=row_idx, column=1, value=lane).font = professional_font
    ws.cell(row=row_idx, column=2, value=loads).font = professional_font
    ws.cell(row=row_idx, column=3, value=avg_rev).font = professional_font
    # Formula for Total Revenue: =B[row]*C[row]
    ws.cell(row=row_idx, column=4, value=f'=B{row_idx}*C{row_idx}').font = professional_font

# 3. Totals Row (Row 6)
ws.cell(row=6, column=1, value='Totals').font = professional_font
# Total Loads: =SUM(B2:B5)
ws.cell(row=6, column=2, value='=SUM(B2:B5)').font = professional_font
# Weighted Average Revenue per Load: =D6/B6
ws.cell(row=6, column=3, value='=D6/B6').font = professional_font
# Total Revenue: =SUM(D2:D5)
ws.cell(row=6, column=4, value='=SUM(D2:D5)').font = professional_font

# Save the workbook
wb.save('shipment_data.xlsx')