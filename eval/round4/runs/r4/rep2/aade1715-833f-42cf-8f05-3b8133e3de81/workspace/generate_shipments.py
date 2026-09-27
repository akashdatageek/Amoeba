import openpyxl
from openpyxl.styles import Font

# Create workbook and sheet
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Shipments"

# Professional font
font_arial = Font(name='Arial')

# Headers
headers = ['Lane', 'Loads', 'Avg Revenue', 'Total Revenue']
ws.append(headers)

# Data
data = [
    ['Lane A', 12, 1850],
    ['Lane B', 7, 2400],
    ['Lane C', 20, 1420],
    ['Lane D', 4, 3100],
]

# Populate data and row formulas
for i, row_data in enumerate(data, start=2):
    ws.cell(row=i, column=1, value=row_data[0])
    ws.cell(row=i, column=2, value=row_data[1])
    ws.cell(row=i, column=3, value=row_data[2])
    # Total Revenue formula for the row
    ws.cell(row=i, column=4, value=f'=B{i}*C{i}')

# Totals Row (Row 6)
ws.cell(row=6, column=1, value='Totals')
ws.cell(row=6, column=2, value='=SUM(B2:B5)')
ws.cell(row=6, column=3, value='=D6/B6')
ws.cell(row=6, column=4, value='=SUM(D2:D5)')

# Apply font to all populated cells
for row in ws.iter_rows(min_row=1, max_row=6, min_col=1, max_col=4):
    for cell in row:
        cell.font = font_arial

# Save file
wb.save('shipment_data.xlsx')
print('File shipment_data.xlsx created successfully.')