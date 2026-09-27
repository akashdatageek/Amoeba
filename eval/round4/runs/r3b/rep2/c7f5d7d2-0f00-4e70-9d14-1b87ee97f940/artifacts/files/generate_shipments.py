import openpyxl
from openpyxl.styles import Font

# Create workbook and sheet
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Shipments"

# Professional font
font_arial = Font(name='Arial')

# Headers
headers = ["Lane", "Loads", "Avg Revenue", "Total Revenue"]
for col_num, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_num)
    cell.value = header
    cell.font = font_arial

# Data
data = [
    ["Lane A", 12, 1850],
    ["Lane B", 7, 2400],
    ["Lane C", 20, 1420],
    ["Lane D", 4, 3100],
]

for row_num, row_data in enumerate(data, 2):
    # Lane, Loads, Avg Revenue
    for col_num, value in enumerate(row_data, 1):
        cell = ws.cell(row=row_num, column=col_num)
        cell.value = value
        cell.font = font_arial
    
    # Total Revenue Formula (B * C)
    total_rev_cell = ws.cell(row=row_num, column=4)
    total_rev_cell.value = f"=B{row_num}*C{row_num}"
    total_rev_cell.font = font_arial

# Totals Row
ws.cell(row=6, column=1).value = "Totals"
ws.cell(row=6, column=1).font = font_arial

# Total Loads Formula
ws.cell(row=6, column=2).value = "=SUM(B2:B5)"
ws.cell(row=6, column=2).font = font_arial

# Total Revenue Formula
ws.cell(row=6, column=4).value = "=SUM(D2:D5)"
ws.cell(row=6, column=4).font = font_arial

# Weighted Average Revenue per Load Formula
ws.cell(row=6, column=3).value = "=D6/B6"
ws.cell(row=6, column=3).font = font_arial

# Formatting
# Column B: Numbers
for row in range(2, 7):
    ws.cell(row=row, column=2).number_format = '#,##0'

# Column C and D: Currency
for row in range(2, 7):
    ws.cell(row=row, column=3).number_format = '"$"#,##0.00'
    ws.cell(row=row, column=4).number_format = '"$"#,##0.00'

# Save file
wb.save("shipment_data.xlsx")