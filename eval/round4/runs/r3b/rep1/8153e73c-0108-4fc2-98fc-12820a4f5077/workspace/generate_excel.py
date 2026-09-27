import openpyxl
from openpyxl.styles import Font

# Create workbook and sheet
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Shipment Data"

# Professional Font
arial_font = Font(name='Arial')

# Headers
headers = ["Lane", "Loads", "Avg Revenue", "Total Revenue", "Weighted Avg Revenue"]
for col_num, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_num)
    cell.value = header
    cell.font = Font(name='Arial', bold=True)

# Data
data = [
    ["Lane A", 12, 1850],
    ["Lane B", 7, 2400],
    ["Lane C", 20, 1420],
    ["Lane D", 4, 3100],
]

# Populate data and row formulas
for row_idx, row_data in enumerate(data, 2):
    ws.cell(row=row_idx, column=1).value = row_data[0] # Lane
    ws.cell(row=row_idx, column=2).value = row_data[1] # Loads
    ws.cell(row=row_idx, column=3).value = row_data[2] # Avg Revenue
    # Total Revenue Formula: Loads * Avg Revenue
    ws.cell(row=row_idx, column=4).value = f"=B{row_idx}*C{row_idx}"

# Totals Row (Row 6)
ws.cell(row=6, column=1).value = "Totals"
ws.cell(row=6, column=1).font = Font(name='Arial', bold=True)
ws.cell(row=6, column=2).value = "=SUM(B2:B5)"
ws.cell(row=6, column=4).value = "=SUM(D2:D5)"
ws.cell(row=6, column=5).value = "=D6/B6"

# Formatting
# Apply Arial to all used cells
for row in ws.iter_rows(max_row=6, max_col=5):
    for cell in row:
        cell.font = arial_font

# Currency formatting for columns C, D, and E
currency_format = '$#,##0'
for row in ws.iter_rows(min_row=2, max_row=6, min_col=3, max_col=5):
    for cell in row:
        cell.number_format = currency_format

# Save the file
wb.save("shipment_data.xlsx")