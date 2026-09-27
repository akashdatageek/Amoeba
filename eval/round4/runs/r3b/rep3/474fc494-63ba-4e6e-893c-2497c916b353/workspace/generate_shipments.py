import openpyxl
from openpyxl.styles import Font

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Shipments"

# Headers based on Step 1 specification
headers = ["Lane", "Loads", "Avg Revenue", "Total Revenue"]
ws.append(headers)

# Data based on Step 1 specification
data = [
    ["Lane A", 12, 1850, 22200],
    ["Lane B", 7, 2400, 16800],
    ["Lane C", 20, 1420, 28400],
    ["Lane D", 4, 3100, 12400],
]

for row in data:
    ws.append(row)

# Totals Row (Row 6) based on Step 1 specification
ws["A6"] = "Totals"
ws["B6"] = "=SUM(B2:B5)"
ws["C6"] = "=D6/B6"
ws["D6"] = "=SUM(D2:D5)"

# Formatting
arial_font = Font(name='Arial')
currency_format = '$#,##0'

for row in ws.iter_rows():
    for cell in row:
        cell.font = arial_font
        if cell.column in [3, 4]: # Avg Revenue and Total Revenue
            cell.number_format = currency_format

wb.save("shipment_data.xlsx")
