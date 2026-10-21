import csv, json, datetime as D
from decimal import Decimal as X, ROUND_HALF_UP
from openpyxl import load_workbook
q=lambda v,p='0.01': v.quantize(X(p),rounding=ROUND_HALF_UP)
rows=list(csv.DictReader(open('eia_prices.csv')))
print('rows',len(rows),rows[0]['week_date'],rows[-1]['week_date'],set(r['series_id'] for r in rows))
ds=[D.date.fromisoformat(r['week_date']) for r in rows]
print('gaps',set((b-a).days for a,b in zip(ds,ds[1:])),'mondays',all(d.weekday()==0 for d in ds))
P=[X(r['price_usd_per_gal']) for r in rows]
print('min/max',min(P),max(P))
S=[q((p-X('1.25'))/X('6.2')) for p in P]
gal=X(35*2300)/X('6.2')
F=[gal*p for p in P]
avg=sum(P)/52
print('avg',avg,'avgsur',sum(S)/52,'sur at avg',q((avg-X('1.25'))/X('6.2')),'annual',q(sum(F)),'avgwk',q(sum(F)/52),'cpm',avg/X('6.2'))
hi=max(range(52),key=lambda i:P[i]);lo=min(range(52),key=lambda i:P[i])
print('hi',ds[hi],P[hi],S[hi],q(F[hi]),'lo',ds[lo],P[lo],S[lo],q(F[lo]))
st=json.load(open('stats.json'))
print({k:st[k] for k in st if k not in('weekly','bands')})
print('weekly len',len(st['weekly']),st['weekly'][0])
bad=0
for w,r,s,f in zip(st['weekly'],rows,S,F):
    ks=list(w.keys())
    vals=[str(v) for v in w.values()]
    if str(s) not in vals and float(s) not in w.values(): bad+=1
print('weekly sur mismatches (loose)',bad)
bb=0
for i,b in enumerate(st['bands']):
    p=X('3.00')+X('0.05')*i
    exp=q((p-X('1.25'))/X('6.2'))
    got=[v for k,v in b.items() if 'sur' in k]
    if not got or X(str(got[0]))!=exp: bb+=1;print('band mismatch',i,b,exp)
print('bands',len(st['bands']),'mismatch',bb,st['bands'][0],st['bands'][-1])
print('atri',st['atri']); print('missing',st['missing'])
wf=load_workbook('fuel_cost_workbook.xlsx'); wv=load_workbook('fuel_cost_workbook.xlsx',data_only=True)
print(wf.sheetnames, wf.defined_names.keys() if hasattr(wf.defined_names,'keys') else wf.defined_names)
for n in wf.sheetnames:
    ws=wf[n]; nf=sum(1 for r in ws.iter_rows() for c in r if isinstance(c.value,str) and c.value.startswith('='))
    print(n,ws.max_row,ws.max_column,'formulas',nf)
I=wf['Inputs']
for r in I.iter_rows(min_row=1,max_row=I.max_row):
    print([ (c.value if not isinstance(c.value,str) else c.value[:60]) for c in r[:3]])
W=wf['Weekly'];Wv=wv['Weekly']
print([c.value for c in W[1]]);print([c.value for c in W[2]]);print([c.value for c in W[53]]);print([c.value for c in W[54]])
mis=0;cnt=0
for i in range(52):
    r=i+2
    pv=Wv.cell(r,2).value;sv=Wv.cell(r,3).value;fv=Wv.cell(r,5).value
    cnt+=1
    if X(str(pv))!=P[i] or sv is None or X(str(round(sv,4)))!=S[i] or abs(fv-float(F[i]))>0.005: mis+=1;print('wk mismatch',r,pv,sv,fv,P[i],S[i],F[i])
print('weekly rows checked',cnt,'mismatch',mis)
Bw=wf['Bands'];Bv=wv['Bands']
for r in (1,2,3,42,43): print([c.value for c in Bw[r]],[c.value for c in Bv[r]])
bm=0
for i in range(41):
    p=X('3.00')+X('0.05')*i
    row=[c.value for c in Bv[i+2]]
    nums=[v for v in row if isinstance(v,(int,float))]
    if not nums or X(str(round(nums[-1],2)))!=q((p-X('1.25'))/X('6.2')): bm+=1
print('Bands sheet mismatch',bm)
for r in wf['Summary'].iter_rows():
    print([c.value for c in r[:3]],[wv['Summary'].cell(r[0].row,2).value])
