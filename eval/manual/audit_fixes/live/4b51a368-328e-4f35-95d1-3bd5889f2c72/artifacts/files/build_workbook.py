import csv, json, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.workbook.defined_name import DefinedName

inp = json.load(open('inputs.json'))
rows = list(csv.DictReader(open('eia_prices.csv', newline='')))
assert len(rows) == 52, len(rows)
SERIES = rows[0]['series_id']; URL = rows[0]['source_url']; RETR = rows[0]['retrieved_date']

F = 'Arial'
bold = Font(name=F, bold=True); norm = Font(name=F); blue = Font(name=F, color='0000FF')
hdr_fill = PatternFill('solid', start_color='D9E1F2')
def hdr(ws, r, vals):
    for i, v in enumerate(vals, 1):
        c = ws.cell(r, i, v); c.font = bold; c.fill = hdr_fill; c.alignment = Alignment(wrap_text=True, vertical='top')

wb = Workbook()
# ---------- Inputs ----------
wi = wb.active; wi.title = 'Inputs'
hdr(wi, 1, ['Input', 'Value', 'Name', 'Source / note'])
items = [
 ('Tractors', inp['tractors'], 'Tractors', 'Given in task (35 Class 8 tractors)'),
 ('Miles per tractor per week', inp['miles_per_tractor_per_week'], 'MilesPerWeek', 'Given in task'),
 ('Fuel economy (mpg)', inp['mpg'], 'MPG', 'Given in task (6.2 mpg)'),
 ('Base (peg) price, $/gal', inp['base_price_per_gal'], 'BasePrice', 'Given in task ($1.25 base)'),
 ('Weeks in window', inp['weeks'], 'Weeks', 'Given in task (52 weeks)'),
 ('Band step, $/gal', inp['band_step'], 'BandStep', 'Given in task (5-cent bands)'),
 ('Band start, $/gal', inp['band_start'], 'BandStart', 'Given in task ($3.00 to $5.00)'),
]
for i, (lab, val, nm, note) in enumerate(items, 2):
    wi.cell(i, 1, lab).font = norm; c = wi.cell(i, 2, val); c.font = blue
    wi.cell(i, 3, nm).font = norm; wi.cell(i, 4, note).font = norm
    wb.defined_names[nm] = DefinedName(nm, attr_text=f"Inputs!$B${i}")
r = len(items) + 2
wi.cell(r, 1, 'Fleet miles per week (derived)').font = norm
wi.cell(r, 2, '=Tractors*MilesPerWeek').font = norm; wi.cell(r, 3, 'FleetMilesWeek').font = norm
wi.cell(r, 4, '35 x 2,300 = 80,500 miles').font = norm
wb.defined_names['FleetMilesWeek'] = DefinedName('FleetMilesWeek', attr_text=f"Inputs!$B${r}")
r += 1
wi.cell(r, 1, 'Fleet gallons per week (derived)').font = norm
wi.cell(r, 2, '=FleetMilesWeek/MPG').font = norm; wi.cell(r, 3, 'FleetGallonsWeek').font = norm
wi.cell(r, 4, '80,500 / 6.2 = 12,983.87 gallons').font = norm
wb.defined_names['FleetGallonsWeek'] = DefinedName('FleetGallonsWeek', attr_text=f"Inputs!$B${r}")
wi.cell(r + 2, 1, 'Legend: blue = typed input (edit here); black = formula. All Weekly/Bands formulas read these named cells.').font = norm
wi.cell(r + 3, 1, 'Price series: EIA weekly retail on-highway diesel, all types, Midwest PADD 2, used as a proxy for Indiana (assumption).').font = norm
for col, w in zip('ABCD', (36, 14, 18, 52)): wi.column_dimensions[col].width = w

# ---------- Weekly ----------
ww = wb.create_sheet('Weekly')
hdr(ww, 1, ['week_date', 'price', 'surcharge_per_mile', 'fleet_gallons', 'fleet_fuel_cost', 'fuel_cost_per_mile'])
for i, row in enumerate(rows, 2):
    d = datetime.datetime.strptime(row['week_date'], '%Y-%m-%d')
    ww.cell(i, 1, d).number_format = 'yyyy-mm-dd'
    ww.cell(i, 2, float(row['price_usd_per_gal'])).number_format = '$0.000'
    ww.cell(i, 2).font = blue
    ww.cell(i, 3, f'=ROUND((B{i}-BasePrice)/MPG,2)').number_format = '$0.00'
    ww.cell(i, 4, '=Tractors*MilesPerWeek/MPG').number_format = '#,##0.00'
    ww.cell(i, 5, f'=D{i}*B{i}').number_format = '$#,##0.00'
    ww.cell(i, 6, f'=B{i}/MPG').number_format = '$0.0000'
    for c in (1, 3, 4, 5, 6): ww.cell(i, c).font = norm
    ww.cell(i, 1).font = blue
first, last = 2, 53
t = last + 1
ww.cell(t, 1, 'TOTAL / AVERAGE').font = bold
ww.cell(t, 2, f'=AVERAGE(B{first}:B{last})').number_format = '$0.000'
ww.cell(t, 3, f'=AVERAGE(C{first}:C{last})').number_format = '$0.0000'
ww.cell(t, 4, f'=SUM(D{first}:D{last})').number_format = '#,##0.00'
ww.cell(t, 5, f'=SUM(E{first}:E{last})').number_format = '$#,##0.00'
ww.cell(t, 6, f'=AVERAGE(F{first}:F{last})').number_format = '$0.0000'
for c in range(2, 7): ww.cell(t, c).font = bold
ww.cell(t + 2, 1, f'Prices: EIA series {SERIES}, Midwest PADD 2 weekly retail on-highway diesel (proxy for Indiana), $/gal; typed from eia_prices.csv, retrieved {RETR}. Row {t}: price, surcharge and fuel cost per mile are averages; gallons and fleet cost are 52-week sums.').font = norm
for col, w in zip('ABCDEF', (16, 10, 20, 16, 18, 18)): ww.column_dimensions[col].width = w
ww.freeze_panes = 'A2'

# ---------- Bands ----------
wbn = wb.create_sheet('Bands')
hdr(wbn, 1, ['lower_bound', 'band_label', 'surcharge_per_mile'])
for k in range(41):
    i = k + 2
    wbn.cell(i, 1, '=BandStart' if k == 0 else f'=ROUND(A{i-1}+BandStep,2)').number_format = '$0.00'
    if k < 40:
        wbn.cell(i, 2, f'="$"&TEXT(A{i},"0.00")&"-"&TEXT(A{i}+BandStep-0.001,"0.000")')
    else:
        wbn.cell(i, 2, '$5.00 and above')
    wbn.cell(i, 3, f'=ROUND((A{i}-BasePrice)/MPG,2)').number_format = '$0.00'
    for c in (1, 2, 3): wbn.cell(i, c).font = norm
wbn.cell(44, 1, 'Surcharge computed at the band lower-bound price: ROUND((price - base)/mpg, 2). Assumption per working interpretation. The 0.001 in the label is a display offset only.').font = norm
for col, w in zip('ABC', (14, 20, 20)): wbn.column_dimensions[col].width = w

# ---------- Summary ----------
ws = wb.create_sheet('Summary')
hdr(ws, 1, ['Metric', 'Value', 'Note / source'])
W = 'Weekly'
S = [
 ('Weeks in window', f'=COUNT({W}!B2:B53)', '0', 'Count of Weekly price rows'),
 ('Window start (week_date)', f'=MIN({W}!A2:A53)', 'yyyy-mm-dd', 'First Monday in Weekly'),
 ('Window end / latest week (week_date)', f'=MAX({W}!A2:A53)', 'yyyy-mm-dd', 'Latest week in the saved EIA page (released 9/22/2026); newer releases not fetched'),
 ('Average price, $/gal', f'=AVERAGE({W}!B2:B53)', '$0.000', f'EIA {SERIES}'),
 ('Highest weekly price, $/gal', f'=MAX({W}!B2:B53)', '$0.000', ''),
 ('Highest week (date)', f'=INDEX({W}!A2:A53,MATCH(B6,{W}!B2:B53,0))', 'yyyy-mm-dd', 'INDEX/MATCH (first match)'),
 ('Lowest weekly price, $/gal', f'=MIN({W}!B2:B53)', '$0.000', ''),
 ('Lowest week (date)', f'=INDEX({W}!A2:A53,MATCH(B8,{W}!B2:B53,0))', 'yyyy-mm-dd', 'INDEX/MATCH (first match)'),
 ('Average surcharge, $/mile', f'=AVERAGE({W}!C2:C53)', '$0.0000', 'Mean of weekly rounded surcharges'),
 ('Highest surcharge, $/mile', f'=MAX({W}!C2:C53)', '$0.00', ''),
 ('Lowest surcharge, $/mile', f'=MIN({W}!C2:C53)', '$0.00', ''),
 ('Annual fleet gallons (52 weeks)', f'=SUM({W}!D2:D53)', '#,##0.00', ''),
 ('Annual fleet fuel cost (52 weeks), $', f'=SUM({W}!E2:E53)', '$#,##0.00', 'No discounts'),
 ('Average fuel cost per mile, $ (carrier, 52 wk)', f'=B14/(FleetMilesWeek*Weeks)', '$0.0000', 'Annual fuel cost / (80,500 miles x 52) = avg price / 6.2'),
 ('ATRI total marginal cost per mile, $ (2025 cost year)', 2.336, '$0.000', 'ATRI An Analysis of the Operational Costs of Trucking: 2026 Update, released 15 July 2026; read from abstract (TRID) https://trid.trb.org/View/2727686 and https://truckingresearch.org/about-atri/atri-research/operational-costs-of-trucking; retrieved 2026-10-16'),
 ('ATRI fuel cost per mile, $ (2025 cost year)', 0.482, '$0.000', 'Same source: 48.2 cents per mile; released 15 July 2026; retrieved 2026-10-16'),
 ('ATRI fuel share of marginal cost', '=B17/B16', '0.0%', 'Derived: 0.482 / 2.336'),
 ('Carrier minus ATRI fuel cost per mile, $', '=B15-B17', '$0.0000', 'Positive = carrier above ATRI fuel benchmark'),
 ('Carrier / ATRI fuel cost per mile', '=B15/B17', '0.00"x"', 'Ratio'),
 ('Carrier fuel cost per mile as share of ATRI total marginal cost', '=B15/B16', '0.0%', 'Comparison only; ATRI total includes non-fuel costs'),
]
for i, (lab, val, fmt, note) in enumerate(S, 2):
    ws.cell(i, 1, lab).font = norm
    c = ws.cell(i, 2, val); c.number_format = fmt
    c.font = blue if isinstance(val, float) else norm
    ws.cell(i, 3, note).font = norm; ws.cell(i, 3).alignment = Alignment(wrap_text=True, vertical='top')
ws.cell(len(S) + 3, 1, 'Caveat: ATRI fuel cost reflects ATRI survey fuel prices and fleet mpg for cost year 2025 (not 6.2 mpg, not Midwest PADD 2 prices); carrier figure covers the 52 weeks to the latest week above.').font = norm
for col, w in zip('ABC', (58, 18, 90)): ws.column_dimensions[col].width = w

# ---------- Sources ----------
wsrc = wb.create_sheet('Sources')
hdr(wsrc, 1, ['Item', 'Detail'])
src = [
 ('EIA series ID', SERIES),
 ('EIA series name', 'Weekly Midwest No 2 Diesel Retail Prices (Dollars per Gallon) - PADD 2 proxy for Indiana (assumption)'),
 ('EIA URL', URL),
 ('EIA retrieved date', RETR),
 ('EIA page release date', '9/22/2026 (saved page); EIA summary page shows later releases 9/29/2026 and 10/6/2026 that were NOT fetched'),
 ('EIA window', f"{rows[0]['week_date']} to {rows[-1]['week_date']} (52 Mondays)"),
 ('Indiana state series', 'MISSING - none seen on the EIA pages read; Midwest PADD 2 used instead'),
 ('ATRI edition', 'An Analysis of the Operational Costs of Trucking: 2026 Update (cost year 2025), released 15 July 2026'),
 ('ATRI URL (report page)', 'https://truckingresearch.org/about-atri/atri-research/operational-costs-of-trucking'),
 ('ATRI URL (abstract)', 'https://trid.trb.org/View/2727686'),
 ('ATRI URL (press release)', 'https://truckingresearch.org/2026/07/new-atri-report-details-accelerating-costs-and-low-profitability-despite-cuts'),
 ('ATRI retrieved date', '2026-10-16'),
 ('ATRI figures', 'Total $2.336/mile; fuel 48.2 cents/mile (2025); 2024: $2.26 and 48.1 cents. Full PDF not read.'),
 ('Fleet inputs', 'Given in the task: 35 tractors, 2,300 miles/week, 6.2 mpg, $1.25 base'),
]
for i, (a, b) in enumerate(src, 2):
    wsrc.cell(i, 1, a).font = bold; wsrc.cell(i, 2, b).font = norm
wsrc.column_dimensions['A'].width = 26; wsrc.column_dimensions['B'].width = 120

wb.save('fuel_cost_workbook.xlsx')
print('saved')
