from openpyxl import load_workbook
from openpyxl.styles import Font
wb=load_workbook('fuel_cost_workbook.xlsx')
S=wb['Summary']; I=wb['Inputs']
r=I.max_row+2
I.cell(r,1,'ATRI inputs (typed once; Summary links here)').font=Font(name='Arial',bold=True)
r+=1
for row in range(1,S.max_row+1):
    c=S.cell(row,2)
    if isinstance(c.value,(int,float)) and not isinstance(c.value,bool):
        lab=S.cell(row,1).value
        I.cell(r,1,lab); x=I.cell(r,2,c.value); x.font=Font(name='Arial',color='0000FF')
        I.cell(r,3,'ATRI Operational Costs of Trucking 2026 Update (cost year 2025), released 15 July 2026; https://truckingresearch.org/about-atri/atri-research/operational-costs-of-trucking')
        print('moved',c.coordinate,lab,c.value,'-> Inputs!B%d'%r)
        c.value='=Inputs!$B$%d'%r; r+=1
wb.move_sheet('Weekly',offset=-wb.sheetnames.index('Weekly'))
wb.active=0
for ws in wb: ws.sheet_view.tabSelected=(ws.title=='Weekly')
wb.save('fuel_cost_workbook.xlsx')
