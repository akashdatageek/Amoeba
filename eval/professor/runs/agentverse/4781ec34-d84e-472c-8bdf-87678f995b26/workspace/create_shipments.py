import openpyxl
from openpyxl.styles import Font

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Shipment Data"

# Professional font
font_arial = Font(name='Arial')

# Headers
headers = ["Lane", "Loads", "Avg Revenue per Load", "Total Revenue"]
ws.append(headers)
for cell in ws[1]:
    cell.font = font_arial

# Data
data = [
    ["Lane A", 12, 1850],
    ["Lane B", 7, 2400],
    ["Lane C", 20, 1420],
    ["Lane D", 4, 3100],
]

# Fill data and add total revenue formulas
for row_idx, row_data in enumerate(data, start=2):
    ws.cell(row=row_idx, column=1, value=row_data[0]).font = font_arial
    ws.cell(row=row_idx, column=2, value=row_data[1]).font = font_arial
    ws.cell(row=row_idx, column=3, value=row_data[2]).font = font_arial
    # Total Revenue = Loads * Avg Revenue
    ws.cell(row=row_idx, column=4, value=f"=B{row_idx}*C{row_idx}").font = font_arial

# Totals Row
totals_row = len(data) + 2
ws.cell(row=totals_row, column=1, value="Total").font = font_arial
ws.cell(row=totals_row, column=2, value=f"=SUM(B2:B{totals_row-1})").font = font_arial

# Average Revenue per Load (Weighted Average: Total Revenue / Total Loads)
ws.cell(row=totals_row, column=3, value=f"=D{totals_row}/B{totals_row}").font = font_arial

# Total Revenue Sum
ws.cell(row=totals_row, column=4, value=f"=SUM(D2:D{totals_row-1})").font = font_arial

wb.save("shipment_data.xlsx")