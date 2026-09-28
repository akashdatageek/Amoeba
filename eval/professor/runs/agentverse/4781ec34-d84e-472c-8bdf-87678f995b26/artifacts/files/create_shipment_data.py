import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Create workbook and sheet
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Shipment Data"

# Data definition
data = [
    ["Lane A", 12, 1850.00],
    ["Lane B", 7, 2400.00],
    ["Lane C", 20, 1420.00],
    ["Lane D", 4, 3100.00],
]

# Headers
headers = ["Lane", "Loads", "Avg Revenue", "Total Revenue", "Vol Contribution %", "Rev Contribution %", "Yield Variance"]
for col_num, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_num, value=header)
    cell.font = Font(bold=True, name="Arial")
    cell.alignment = Alignment(horizontal="center")

# Populate data and formulas
for row_num, row_data in enumerate(data, 2):
    # Basic data
    ws.cell(row=row_num, column=1, value=row_data[0])
    ws.cell(row=row_num, column=2, value=row_data[1])
    ws.cell(row=row_num, column=3, value=row_data[2])
    
    # Total Revenue: =B2*C2
    ws.cell(row=row_num, column=4, value=f"=B{row_num}*C{row_num}")
    
    # Vol Contribution %: =B2/$B$6
    ws.cell(row=row_num, column=5, value=f"=B{row_num}/$B$6")
    
    # Rev Contribution %: =D2/$D$6
    ws.cell(row=row_num, column=6, value=f"=D{row_num}/$D$6")
    
    # Yield Variance: =C2-$C$6
    ws.cell(row=row_num, column=7, value=f"=C{row_num}-$C$6")

# Totals Row (Row 6)
totals_row = 6
ws.cell(row=totals_row, column=1, value="Total").font = Font(bold=True, name="Arial")

# Total Loads: =SUM(B2:B5)
ws.cell(row=totals_row, column=2, value="=SUM(B2:B5)").font = Font(bold=True, name="Arial")

# Weighted Avg Revenue: =IFERROR(D6/B6, 0)
ws.cell(row=totals_row, column=3, value=f"=IFERROR(D{totals_row}/B{totals_row}, 0)").font = Font(bold=True, name="Arial")

# Total Revenue: =SUM(D2:D5)
ws.cell(row=totals_row, column=4, value=f"=SUM(D2:D{totals_row-1})").font = Font(bold=True, name="Arial")

# Formatting
thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

for row in ws.iter_rows(min_row=1, max_row=totals_row, min_col=1, max_col=7):
    for cell in row:
        cell.font = Font(name="Arial")
        cell.border = thin_border
        
        # Alignment
        if cell.column == 1: # Lane
            cell.alignment = Alignment(horizontal="left")
        elif cell.column == 2: # Loads
            cell.alignment = Alignment(horizontal="center")
        else: # Revenue and Percentages
            cell.alignment = Alignment(horizontal="right")
        
        # Number Formatting
        if cell.column in [3, 4, 7]: # Currency
            cell.number_format = '$#,##0.00'
        elif cell.column in [5, 6]: # Percentage
            cell.number_format = '0.00%'

# Adjust column widths for readability
column_widths = [15, 12, 15, 15, 20, 20, 15]
for i, width in enumerate(column_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = width

wb.save("shipment_data.xlsx")
