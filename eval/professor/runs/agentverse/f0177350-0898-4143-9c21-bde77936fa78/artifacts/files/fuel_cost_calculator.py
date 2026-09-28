import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

def create_fuel_cost_sheet():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fuel Cost Analysis"

    # --- Data & Parameters ---
    diesel_price = 5.946  # EIA Midwest (PADD 2) as of 09/07/26
    fuel_efficiency = 6.5
    contingency_factor = 0.10
    
    legs = [
        ("Chicago–Indianapolis", 185),
        ("Indianapolis–Columbus", 175),
        ("Columbus–Pittsburgh", 185),
    ]

    # --- Trip Parameters Block ---
    ws['G2'] = "Trip Parameters"
    ws['G3'] = "Diesel Price (Midwest)"
    ws['H3'] = diesel_price
    ws['G4'] = "Fuel Efficiency (mpg)"
    ws['H4'] = fuel_efficiency
    ws['G5'] = "Contingency Factor"
    ws['H5'] = contingency_factor

    # Define Named Ranges
    wb.create_named_range('DieselPrice', ws, '$H$3')
    wb.create_named_range('FuelEfficiency', ws, '$H$4')
    wb.create_named_range('ContingencyFactor', ws, '$H$5')

    # --- Main Table ---
    headers = ["Leg", "Distance (mi)", "Gallons", "Cost"]
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    # Data Rows
    for row_idx, (leg_name, distance) in enumerate(legs, 2):
        ws.cell(row=row_idx, column=1).value = leg_name
        ws.cell(row=row_idx, column=2).value = distance
        ws.cell(row=row_idx, column=3).value = f"=B{row_idx}/FuelEfficiency"
        ws.cell(row=row_idx, column=4).value = f"=C{row_idx}*DieselPrice"

    # Total Row
    total_row = len(legs) + 2
    ws.cell(row=total_row, column=1).value = "Total"
    ws.cell(row=total_row, column=1).font = Font(bold=True)
    ws.cell(row=total_row, column=2).value = f"=SUM(B2:B{total_row-1})"
    ws.cell(row=total_row, column=3).value = f"=SUM(C2:C{total_row-1})"
    ws.cell(row=total_row, column=4).value = f"=SUM(D2:D{total_row-1})"
    ws.cell(row=total_row, column=2).font = Font(bold=True)
    ws.cell(row=total_row, column=3).font = Font(bold=True)
    ws.cell(row=total_row, column=4).font = Font(bold=True)

    # Contingency Row
    cont_row = total_row + 1
    ws.cell(row=cont_row, column=1).value = "Contingency (10%)"
    ws.cell(row=cont_row, column=4).value = f"=D{total_row}*ContingencyFactor"

    # Grand Total Row
    grand_total_row = cont_row + 1
    ws.cell(row=grand_total_row, column=1).value = "Grand Total"
    ws.cell(row=grand_total_row, column=1).font = Font(bold=True)
    ws.cell(row=grand_total_row, column=4).value = f"=D{total_row}+D{cont_row}"
    ws.cell(row=grand_total_row, column=4).font = Font(bold=True)

    # --- Formatting ---
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name="Arial")

    for row in range(2, grand_total_row + 1):
        ws.cell(row=row, column=3).number_format = '0.00'
        ws.cell(row=row, column=4).number_format = '$#,##0.00'
    
    ws['H3'].number_format = '$#,##0.00'
    ws['H5'].number_format = '0%'

    thin = Side(border_style="thin", color="000000")
    border = Border(top=thin, left=thin, right=thin, bottom=thin)
    for row in ws.iter_rows(min_row=1, max_row=grand_total_row, min_col=1, max_col=4):
        for cell in row:
            cell.border = border

    ws['G2'].font = Font(bold=True)
    for r in range(3, 6):
        ws.cell(row=r, column=7).font = Font(italic=True)

    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['G'].width = 25

    ws['G7'] = "Source: EIA Weekly Retail Gasoline and Diesel Prices (Midwest PADD 2), 09/07/26"
    ws['G7'].font = Font(size=9, italic=True)

    wb.save("Fuel_Cost_Analysis.xlsx")

if __name__ == "__main__":
    create_fuel_cost_sheet()