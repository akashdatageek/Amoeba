import openpyxl
from openpyxl.styles import Font

# Create workbook and sheet
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Shipments"

# Professional font
arial_font = Font(name='Arial')

# Headers
headers = ["Lane", "Loads", "Avg Revenue", "Total Revenue"]
ws.append(headers)
for cell in ws[1]:
    cell.font = arial_font

# Data
data = [
    ["Lane A", 12, 1850],
    ["Lane B", 7, 2400],
    ["Lane C", 20, 1420],
    ["Lane D", 4, 3100],
]

# Populate data and row formulas
for i, row_data in enumerate(data, start=2):
    ws.cell(row=i, column=1, value=row_data[0]).font = arial_font
    ws.cell(row=i, column=2, value=row_data[1]).font = arial_font
    ws.cell(row=i, column=3, value=row_data[2]).font = arial_font
    # Total Revenue formula: Loads * Avg Revenue
    ws.cell(row=i, column=4, value=f"=B{i}*C{i}").font = arial_font

# Totals row (Row 6)
ws.cell(row=6, column=1, value="Total").font = arial_font
ws.cell(row=6, column=2, value="=SUM(B2:B5)").font = arial_font
ws.cell(row=6, column=3, value="=SUM(D2:D5)/SUM(B2:B5)").font = arial_font
ws.cell(row=6, column=4, value="=SUM(D2:D5)").font = arial_font

# Save the file
wb.save("shipment_data.xlsx")