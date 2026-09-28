import openpyxl
from openpyxl.styles import Font, Alignment

# Create workbook and sheet
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Fuel Costs"

# Define headers and data
headers = ["Leg", "Distance (mi)", "Gallons", "Cost (USD)"]
data = [
    ["Chicago-Indy", 185, 28.46, 185.83],
    ["Indy-Columbus", 175, 26.92, 175.77],
    ["Columbus-Pitt", 185, 28.46, 185.83],
]

# Write headers
for col_num, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_num, value=header)
    cell.font = Font(name='Arial', bold=True)
    cell.alignment = Alignment(horizontal='center')

# Write data
for row_num, row_data in enumerate(data, 2):
    for col_num, value in enumerate(row_data, 1):
        cell = ws.cell(row=row_num, column=col_num, value=value)
        cell.font = Font(name='Arial')

# Add Total row
total_row = len(data) + 2
ws.cell(row=total_row, column=1, value="Total").font = Font(name='Arial', bold=True)

# Distance Total (Optional but professional)
ws.cell(row=total_row, column=2, value=f"=SUM(B2:B{total_row-1})").font = Font(name='Arial', bold=True)

# Gallons Total (Required Formula)
ws.cell(row=total_row, column=3, value=f"=SUM(C2:C{total_row-1})").font = Font(name='Arial', bold=True)

# Cost Total (Required Formula)
ws.cell(row=total_row, column=4, value=f"=SUM(D2:D{total_row-1})").font = Font(name='Arial', bold=True)

# Adjust column widths for professionalism
ws.column_dimensions['A'].width = 20
ws.column_dimensions['B'].width = 15
ws.column_dimensions['C'].width = 12
ws.column_dimensions['D'].width = 12

# Save file
wb.save("fuel_cost_file.xlsx")
