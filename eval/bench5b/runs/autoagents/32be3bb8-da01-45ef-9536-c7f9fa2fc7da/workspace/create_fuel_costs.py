import openpyxl
from openpyxl.styles import Font, Alignment

def create_excel():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fuel Costs"

    # Define headers
    headers = ['Leg', 'Gallons', 'Cost']
    ws.append(headers)

    # Define data
    data = [
        ['Chicago–Indianapolis', 28.46, 185.83],
        ['Indianapolis–Columbus', 26.92, 175.78],
        ['Columbus–Pittsburgh', 28.46, 185.83],
    ]

    # Populate data
    for row in data:
        ws.append(row)

    # Add Total row
    # Row 1: Header, Row 2-4: Data, Row 5: Total
    ws.append(['Total', '=SUM(B2:B4)', '=SUM(C2:C4)'])

    # Formatting
    header_font = Font(bold=True, name='Arial')
    total_font = Font(bold=True, name='Arial')
    
    # Apply header formatting
    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')

    # Apply total row formatting
    ws['A5'].font = total_font
    ws['B5'].font = total_font
    ws['C5'].font = total_font

    # Set column widths for professionalism
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 12

    # Save the file
    wb.save('fuel_costs.xlsx')
    print("fuel_costs.xlsx created successfully.")

if __name__ == "__main__":
    create_excel()
