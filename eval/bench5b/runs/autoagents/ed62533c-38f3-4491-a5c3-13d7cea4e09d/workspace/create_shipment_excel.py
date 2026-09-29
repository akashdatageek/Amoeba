import openpyxl
from openpyxl.styles import Font

# Data
data = [
    ["Lane A", 12, 1850],
    ["Lane B", 7, 2400],
    ["Lane C", 20, 1420],
    ["Lane D", 4, 3100],
]

# Create workbook and sheet
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Shipment Data"

# Headers
headers = ['Lane', 'Loads', 'Avg Revenue per Load', 'Total Revenue']
ws.append(headers)

# Professional Font
font_arial = Font(name='Arial')

# Populate data and add Total Revenue formulas
for row_idx, row_data in enumerate(data, start=2):
    ws.cell(row=row_idx, column=1, value=row_data[0])
    ws.cell(row=row_idx, column=2, value=row_data[1])
    ws.cell(row=row_idx, column=3, value=row_data[2])
    # Total Revenue = Loads * Avg Revenue
    ws.cell(row=row_idx, column=4, value=f"=B{row_idx}*C{row_idx}")

# Totals Row
totals_row = len(data) + 2
ws.cell(row=totals_row, column=1, value="Totals")
ws.cell(row=totals_row, column=2, value=f"=SUM(B2:B{totals_row-1})")
ws.cell(row=totals_row, column=4, value=f"=SUM(D2:D{totals_row-1})")

# Weighted Average Row
avg_row = totals_row + 1
ws.cell(row=avg_row, column=1, value="Average Revenue per Load")
# Formula: Total Revenue Sum / Total Loads Sum
ws.cell(row=avg_row, column=2, value=f"=D{totals_row}/B{totals_row}")

# Formatting
currency_format = '$#,##0.00'

for row in ws.iter_rows():
    for cell in row:
        cell.font = font_arial
        # Apply currency format to columns C and D, and the final average cell
        if cell.column == 3 or cell.column == 4:
            cell.number_format = currency_format
        if cell.row == avg_row and cell.column == 2:
            cell.number_format = currency_format

# Save the file
wb.save('shipment_data.xlsx')
print("File shipment_data.xlsx created successfully.")
