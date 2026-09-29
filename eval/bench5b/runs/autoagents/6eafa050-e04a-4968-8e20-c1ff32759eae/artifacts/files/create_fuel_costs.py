import openpyxl
from openpyxl.styles import Font

# Create a new workbook and select the active sheet
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Fuel Costs"

# Define the header and data
headers = ['Leg', 'Gallons', 'Cost']
data = [
    ['Chicago-Indy', 28.46, 184.27],
    ['Indy-Columbus', 26.92, 174.22],
    ['Columbus-Pitt', 28.46, 184.27],
]

# 1. Create header row
ws.append(headers)

# 2. Populate the three legs
for row in data:
    ws.append(row)

# 3. Add a 'Total' row
# Row 1 is header, Rows 2, 3, 4 are data. Total is Row 5.
total_row_idx = 5
ws.cell(row=total_row_idx, column=1, value='Total')

# 4. Insert actual Excel formulas for Gallons and Cost
ws.cell(row=total_row_idx, column=2, value='=SUM(B2:B4)')
ws.cell(row=total_row_idx, column=3, value='=SUM(C2:C4)')

# Professional formatting: Set font to Arial for all used cells
arial_font = Font(name='Arial')
for row in ws.iter_rows():
    for cell in row:
        cell.font = arial_font

# Save the file
wb.save('fuel_costs.xlsx')
print("fuel_costs.xlsx created successfully.")
