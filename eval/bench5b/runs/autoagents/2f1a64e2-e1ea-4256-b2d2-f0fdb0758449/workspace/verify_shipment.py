import openpyxl

def verify():
    wb = openpyxl.load_workbook('shipment_data.xlsx', data_only=False)
    sheet = wb.active
    
    # Expected Data
    expected_data = {
        "Lane A": {"loads": 12, "avg": 1850},
        "Lane B": {"loads": 7, "avg": 2400},
        "Lane C": {"loads": 20, "avg": 1420},
        "Lane D": {"loads": 4, "avg": 3100},
    }
    
    results = []
    
    # We assume data starts from row 2, columns: Lane, Loads, Avg Revenue, Total Revenue
    # Row 1: Header
    # Row 2-5: Lanes
    # Row 6: Totals
    
    # Check Lane Data
    for i, (lane, vals) in enumerate(expected_data.items(), start=2):
        cell_lane = sheet.cell(row=i, column=1).value
        cell_loads = sheet.cell(row=i, column=2).value
        cell_avg = sheet.cell(row=i, column=3).value
        cell_total = sheet.cell(row=i, column=4).value
        
        # Check input data
        if cell_lane != lane or cell_loads != vals['loads'] or cell_avg != vals['avg']:
            results.append(f"Mismatch in {lane}: Expected {vals}, found {cell_loads}, {cell_avg}")
        
        # Check Total Revenue calculation (Manual check)
        expected_total = vals['loads'] * vals['avg']
        # If cell_total is a formula, we can't easily evaluate it without data_only=True, 
        # but we can check if it's a formula or a value.
        # To check the value, we need another load.
        
    # Load with data_only=True to check calculated values
    wb_val = openpyxl.load_workbook('shipment_data.xlsx', data_only=True)
    sheet_val = wb_val.active
    
    for i, (lane, vals) in enumerate(expected_data.items(), start=2):
        val_total = sheet_val.cell(row=i, column=4).value
        expected_total = vals['loads'] * vals['avg']
        if val_total != expected_total:
            results.append(f"Total Revenue mismatch for {lane}: Expected {expected_total}, found {val_total}")

    # Check Totals Row (Row 6)
    total_loads = sheet_val.cell(row=6, column=2).value
    total_revenue = sheet_val.cell(row=6, column=4).value
    
    if total_loads != 43:
        results.append(f"Total Loads mismatch: Expected 43, found {total_loads}")
    if total_revenue != 79800:
        results.append(f"Total Revenue mismatch: Expected 79800, found {total_revenue}")
        
    # Check Average Revenue Formula (Row 6, Column 3)
    # Formula should be =D6/B6 (Total Revenue / Total Loads)
    formula_avg = sheet.cell(row=6, column=3).value
    if not (isinstance(formula_avg, str) and 'D6' in formula_avg and 'B6' in formula_avg):
        results.append(f"Average Revenue formula is incorrect: {formula_avg}")
    else:
        results.append("Average Revenue formula is correct (Weighted Average).")

    if not results or (len(results) == 1 and "correct" in results[0]):
        print("VERIFICATION_SUCCESS")
    else:
        print("VERIFICATION_FAILED")
        for res in results:
            print(res)

verify()
