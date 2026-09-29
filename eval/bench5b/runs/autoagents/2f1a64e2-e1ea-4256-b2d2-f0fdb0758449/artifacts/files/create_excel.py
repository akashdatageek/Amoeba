import openpyxl
from openpyxl.styles import Font

# Data
data = [
    ("Lane A", 12, 1850),
    ("Lane B", 7, 2400),
    ("Lane C", 20, 1420),
    ("Lane D", 4, 3100),
]

# Create workbook and sheet
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Shipment Data"

# Headers
headers = ['Lane', 'Loads', 'Avg Revenue per Load', 'Total Revenue']
ws.append(headers)

# Professional Font
arial_font = Font(name='Arial')

# Write data and formulas for Total Revenue
for i, row_data in enumerate(data, start=2):
    lane, loads, avg_rev = row_data
    ws.cell(row=i, column=1, value=lane)
    ws.cell(row=i, column=2, value=loads)
    ws.cell(row=i, column=3, value=avg_rev)
    # Formula: Loads * Avg Revenue
    ws.cell(row=i, column=4, value=f"=B{i}*C{i}")

# Totals Row
totals_row = len(data) + 2
ws.cell(row=totals_row, column=1, value="Totals")
ws.cell(row=totals_row, column=2, value=f"=SUM(B2:B{totals_row-1})")
ws.cell(row=totals_row, column=4, value=f"=SUM(D2:D{totals_row-1})")

# Weighted Average Revenue per Load
avg_row = totals_row + 1
ws.cell(row=avg_row, column=1, value="Average Revenue per Load")
# Formula: Total Revenue Sum / Total Loads Sum
ws.cell(row=avg_row, column=2, value=f"=D{totals_row}/B{totals_row}")

# Formatting
currency_format = '"$"#,##0'

for row in ws.iter_rows(min_row=1, max_row=avg_row, min_col=1, max_col=4):
    for cell in row:
        cell.font = arial_font
        # Apply currency format to Avg Revenue and Total Revenue columns
        if cell.column == 3 or cell.column == 4:
            cell.number_format = currency_format
        # Special case for the weighted average result cell
        if cell.row == avg_row and cell.column == 2:
            cell.number_format = currency_format

# Save the file
wb.save('shipment_data.xlsx')
