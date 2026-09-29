import openpyxl

def verify():
    # Load workbook without data_only to see formulas
    wb = openpyxl.load_workbook('shipment_data.xlsx', data_only=False)
    ws = wb.active
    
    results = []
    errors = []

    # 1. Verify Input Data
    expected_data = {
        2: {'lane': 'Lane A', 'loads': 12, 'avg': 1850},
        3: {'lane': 'Lane B', 'loads': 7, 'avg': 2400},
        4: {'lane': 'Lane C', 'loads': 20, 'avg': 1420},
        5: {'lane': 'Lane D', 'loads': 4, 'avg': 3100},
    }
    
    for row, vals in expected_data.items():
        lane = ws.cell(row=row, column=1).value
        loads = ws.cell(row=row, column=2).value
        avg = ws.cell(row=row, column=3).value
        if lane == vals['lane'] and loads == vals['loads'] and avg == vals['avg']:
            results.append(f"[PASS] Row {row} ({lane}): Data matches.")
        else:
            errors.append(f"[FAIL] Row {row} ({lane}): Data mismatch. Found {lane}, {loads}, {avg}")

    # 2. Verify Lane Total Revenue Formulas
    for row in range(2, 6):
        formula = ws.cell(row=row, column=4).value
        expected_formula = f'=B{row}*C{row}'
        if formula == expected_formula:
            results.append(f"[PASS] Row {row}: Total Revenue formula matches ({expected_formula}).")
        else:
            errors.append(f"[FAIL] Row {row}: Total Revenue formula mismatch. Expected {expected_formula}, found {formula}")

    # 3. Verify Grand Totals Formulas
    # Total Loads (B6)
    loads_total_formula = ws.cell(row=6, column=2).value
    if loads_total_formula == '=SUM(B2:B5)':
        results.append("[PASS] Grand Total Loads: Formula matches (=SUM(B2:B5)).")
    else:
        errors.append(f"[FAIL] Grand Total Loads: Formula mismatch. Expected =SUM(B2:B5), found {loads_total_formula}")

    # Total Revenue (D6)
    rev_total_formula = ws.cell(row=6, column=4).value
    if rev_total_formula == '=SUM(D2:D5)':
        results.append("[PASS] Grand Total Revenue: Formula matches (=SUM(D2:D5)).")
    else:
        errors.append(f"[FAIL] Grand Total Revenue: Formula mismatch. Expected =SUM(D2:D5), found {rev_total_formula}")

    # 4. Verify Weighted Average Formula (B7)
    avg_formula = ws.cell(row=7, column=2).value
    expected_avg_formula = '=SUM(D2:D5)/SUM(B2:B5)'
    if avg_formula == expected_avg_formula:
        results.append(f"[PASS] Weighted Average: Formula matches ({expected_avg_formula}).")
    else:
        errors.append(f"[FAIL] Weighted Average: Formula mismatch. Expected {expected_avg_formula}, found {avg_formula}")

    # Print results
    print("\n".join(results))
    if errors:
        print("\nERRORS FOUND:")
        print("\n".join(errors))
        print("\nVERIFICATION_FAILED")
    else:
        print("\nVERIFICATION_SUCCESSFUL")

if __name__ == "__main__":
    verify()
