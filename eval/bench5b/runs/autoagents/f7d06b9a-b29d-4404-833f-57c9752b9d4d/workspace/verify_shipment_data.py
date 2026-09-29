import openpyxl

def verify():
    try:
        wb = openpyxl.load_workbook('shipment_data.xlsx')
        ws = wb.active
        
        results = []
        errors = []

        # 1. Verify Input Data
        expected_data = {
            2: {'Lane': 'Lane A', 'Loads': 12, 'Avg': 1850},
            3: {'Lane': 'Lane B', 'Loads': 7, 'Avg': 2400},
            4: {'Lane': 'Lane C', 'Loads': 20, 'Avg': 1420},
            5: {'Lane': 'Lane D', 'Loads': 4, 'Avg': 3100},
        }
        
        for row, vals in expected_data.items():
            lane = ws.cell(row=row, column=1).value
            loads = ws.cell(row=row, column=2).value
            avg = ws.cell(row=row, column=3).value
            if lane == vals['Lane'] and loads == vals['Loads'] and avg == vals['Avg']:
                results.append(f"[PASS] Row {row} ({lane}): Data matches.")
            else:
                errors.append(f"[FAIL] Row {row} ({lane}): Data mismatch. Expected {vals}, found {lane}, {loads}, {avg}")

        # 2. Verify Lane Total Revenue Formulas
        for row in range(2, 6):
            formula = ws.cell(row=row, column=4).value
            expected_formula = f'=B{row}*C{row}'
            if formula == expected_formula:
                results.append(f"[PASS] Row {row}: Total Revenue formula matches ({expected_formula}).")
            else:
                errors.append(f"[FAIL] Row {row}: Total Revenue formula mismatch. Expected {expected_formula}, found {formula}")

        # 3. Verify Grand Totals
        # Total Loads (Row 6, Col 2)
        total_loads_formula = ws.cell(row=6, column=2).value
        if total_loads_formula == '=SUM(B2:B5)':
            results.append("[PASS] Grand Total Loads: Formula matches (=SUM(B2:B5)).")
        else:
            errors.append(f"[FAIL] Grand Total Loads: Formula mismatch. Expected =SUM(B2:B5), found {total_loads_formula}")

        # Total Revenue (Row 6, Col 4)
        total_rev_formula = ws.cell(row=6, column=4).value
        if total_rev_formula == '=SUM(D2:D5)':
            results.append("[PASS] Grand Total Revenue: Formula matches (=SUM(D2:D5)).")
        else:
            errors.append(f"[FAIL] Grand Total Revenue: Formula mismatch. Expected =SUM(D2:D5), found {total_rev_formula}")

        # 4. Verify Weighted Average Formula (Row 7, Col 2)
        avg_rev_formula = ws.cell(row=7, column=2).value
        if avg_rev_formula == '=SUM(D2:D5)/SUM(B2:B5)':
            results.append("[PASS] Weighted Average: Formula matches (=SUM(D2:D5)/SUM(B2:B5)).")
        else:
            errors.append(f"[FAIL] Weighted Average: Formula mismatch. Expected =SUM(D2:D5)/SUM(B2:B5), found {avg_rev_formula}")

        # Print all results
        for r in results:
            print(r)
        for e in errors:
            print(e)
            
        if not errors:
            print("\nVERIFICATION_SUCCESS")
        else:
            print("\nVERIFICATION_FAILED")

    except Exception as e:
        print(f"Error during verification: {e}")

if __name__ == "__main__":
    verify()
