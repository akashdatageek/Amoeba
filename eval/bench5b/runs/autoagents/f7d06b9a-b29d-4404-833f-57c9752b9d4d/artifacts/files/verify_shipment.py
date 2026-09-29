import openpyxl

def verify():
    try:
        wb = openpyxl.load_workbook('shipment_data.xlsx', data_only=False)
        ws = wb.active
        
        results = []
        errors = []
        
        # Expected Data
        expected_data = {
            2: {'lane': 'Lane A', 'loads': 12, 'avg': 1850, 'formula': '=B2*C2'},
            3: {'lane': 'Lane B', 'loads': 7, 'avg': 2400, 'formula': '=B3*C3'},
            4: {'lane': 'Lane C', 'loads': 20, 'avg': 1420, 'formula': '=B4*C4'},
            5: {'lane': 'Lane D', 'loads': 4, 'avg': 3100, 'formula': '=B5*C5'},
        }
        
        # 1. Verify Lane Data and Total Revenue Formulas
        for row, exp in expected_data.items():
            lane = ws.cell(row=row, column=1).value
            loads = ws.cell(row=row, column=2).value
            avg = ws.cell(row=row, column=3).value
            total_formula = ws.cell(row=row, column=4).value
            
            if lane == exp['lane'] and loads == exp['loads'] and avg == exp['avg'] and total_formula == exp['formula']:
                results.append(f"[PASS] Row {row} ({lane}): Data and formula match.")
            else:
                errors.append(f"[FAIL] Row {row} ({lane}): Expected loads={exp['loads']}, avg={exp['avg']}, formula={exp['formula']}. Found loads={loads}, avg={avg}, formula={total_formula}")

        # 2. Verify Grand Totals
        total_loads_formula = ws.cell(row=6, column=2).value
        total_rev_formula = ws.cell(row=6, column=4).value
        
        if total_loads_formula == '=SUM(B2:B5)':
            results.append("[PASS] Grand Total Loads: Formula matches (=SUM(B2:B5)).")
        else:
            errors.append(f"[FAIL] Grand Total Loads: Expected =SUM(B2:B5), found {total_loads_formula}")
            
        if total_rev_formula == '=SUM(D2:D5)':
            results.append("[PASS] Grand Total Revenue: Formula matches (=SUM(D2:D5)).")
        else:
            errors.append(f"[FAIL] Grand Total Revenue: Expected =SUM(D2:D5), found {total_rev_formula}")

        # 3. Verify Weighted Average Formula
        # Based on the Data Engineer's script, it should be in cell B7 (row 7, col 2)
        avg_rev_formula = ws.cell(row=7, column=2).value
        if avg_rev_formula == '=SUM(D2:D5)/SUM(B2:B5)':
            results.append("[PASS] Weighted Average: Formula matches (=SUM(D2:D5)/SUM(B2:B5)).")
        else:
            errors.append(f"[FAIL] Weighted Average: Expected =SUM(D2:D5)/SUM(B2:B5), found {avg_rev_formula}")

        # Print results
        print("\n".join(results))
        if errors:
            print("\nERRORS FOUND:")
            print("\n".join(errors))
            print("\nVERIFICATION_FAILED")
        else:
            print("\nVERIFICATION_SUCCESS")

    except Exception as e:
        print(f"Error during verification: {e}")

if __name__ == "__main__":
    verify()
