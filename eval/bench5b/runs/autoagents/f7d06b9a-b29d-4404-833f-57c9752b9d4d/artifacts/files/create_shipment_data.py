from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

def create_excel():
    wb = Workbook()
    ws = wb.active
    ws.title = "Shipment Data"

    # Headers
    headers = ['Lane', 'Loads', 'Avg Revenue per Load', 'Total Revenue']
    ws.append(headers)

    # Data
    data = [
        ['Lane A', 12, 1850],
        ['Lane B', 7, 2400],
        ['Lane C', 20, 1420],
        ['Lane D', 4, 3100],
    ]

    # Add data and lane total formulas
    for i, row_data in enumerate(data, start=2):
        ws.cell(row=i, column=1, value=row_data[0])
        ws.cell(row=i, column=2, value=row_data[1])
        ws.cell(row=i, column=3, value=row_data[2])
        # Total Revenue = Loads * Avg Revenue
        ws.cell(row=i, column=4, value=f'=B{i}*C{i}')

    # Totals Row
    totals_row = 6
    ws.cell(row=totals_row, column=1, value='Totals')
    ws.cell(row=totals_row, column=2, value='=SUM(B2:B5)')
    ws.cell(row=totals_row, column=4, value='=SUM(D2:D5)')

    # Weighted Average Row
    avg_row = 7
    ws.cell(row=avg_row, column=1, value='Average Revenue per Load')
    # Formula: Total Revenue sum / Total Loads sum
    ws.cell(row=avg_row, column=2, value='=SUM(D2:D5)/SUM(B2:B5)')

    # Formatting
    # Professional font (Arial)
    font_bold = Font(bold=True, name='Arial')
    font_regular = Font(name='Arial')

    # Apply font to all cells
    for row in ws.iter_rows():
        for cell in row:
            cell.font = font_regular

    # Header font and alignment
    for cell in ws[1]:
        cell.font = font_bold
        cell.alignment = Alignment(horizontal='center')

    # Totals and Average row font
    ws.cell(row=totals_row, column=1).font = font_bold
    ws.cell(row=avg_row, column=1).font = font_bold

    # Currency formatting for columns C and D, and the final average
    currency_format = '$#,##0.00'
    for row in range(2, totals_row + 1):
        ws.cell(row=row, column=3).number_format = currency_format
        ws.cell(row=row, column=4).number_format = currency_format

    ws.cell(row=avg_row, column=2).number_format = currency_format

    # Column widths for readability
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 20

    # Save the file
    wb.save('shipment_data.xlsx')
    print("File shipment_data.xlsx created successfully.")

if __name__ == "__main__":
    create_excel()
