from openpyxl import Workbook
from openpyxl.styles import Font

def create_fuel_costs_excel():
    # Create a new workbook and select the active sheet
    wb = Workbook()
    ws = wb.active
    ws.title = "Fuel Costs"

    # 1. Create a header row
    headers = ['Leg', 'Gallons', 'Cost']
    ws.append(headers)

    # 2. Populate the three legs of the trip
    data = [
        ['Chicago-Indy', 28.46, 184.27],
        ['Indy-Columbus', 26.92, 174.22],
        ['Columbus-Pitt', 28.46, 184.27],
    ]
    for row in data:
        ws.append(row)

    # 3. Add a 'Total' row at the bottom
    # The data starts at row 2, ends at row 4. Total is row 5.
    total_row_idx = 5
    ws.cell(row=total_row_idx, column=1, value='Total')

    # 4. Use openpyxl to insert actual Excel formulas
    ws.cell(row=total_row_idx, column=2, value='=SUM(B2:B4)')
    ws.cell(row=total_row_idx, column=3, value='=SUM(C2:C4)')

    # Professional font (Arial) throughout
    arial_font = Font(name='Arial')
    for row in ws.iter_rows():
        for cell in row:
            cell.font = arial_font

    # 5. Save the file to the workspace
    wb.save('fuel_costs.xlsx')
    print("fuel_costs.xlsx created successfully.")

if __name__ == "__main__":
    create_fuel_costs_excel()
