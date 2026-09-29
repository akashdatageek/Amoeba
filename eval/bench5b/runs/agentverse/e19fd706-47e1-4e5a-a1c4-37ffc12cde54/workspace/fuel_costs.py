import openpyxl
from openpyxl.styles import Font

# Data
diesel_price = 4.4798
mpg = 6.5
legs = [
    ("Chicago–Indianapolis", 185),
    ("Indianapolis–Columbus", 175),
    ("Columbus–Pittsburgh", 185),
]
source = "AAA (https://gasprices.aaa.com)"

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Fuel Costs"

# Professional Font
arial_font = Font(name='Arial')
bold_arial = Font(name='Arial', bold=True)

# Setup Constants
ws['A1'] = "Diesel Price (per gallon)"
ws['B1'] = diesel_price
ws['A2'] = "Fuel Efficiency (MPG)"
ws['B2'] = mpg
ws['A3'] = "Source"
ws['B3'] = source

# Table Headers
headers = ["Leg", "Distance (miles)", "Gallons", "Cost"]
for col, text in enumerate(headers, 1):
    cell = ws.cell(row=5, column=col, value=text)
    cell.font = bold_arial

# Data Rows
for i, (leg, dist) in enumerate(legs, 6):
    ws.cell(row=i, column=1, value=leg).font = arial_font
    ws.cell(row=i, column=2, value=dist).font = arial_font
    # Gallons = Distance / MPG
    ws.cell(row=i, column=3, value=f"=B{i}/$B$2").font = arial_font
    # Cost = Gallons * Price
    ws.cell(row=i, column=4, value=f"=C{i}*$B$1").font = arial_font

# Total Row
total_row = 6 + len(legs)
ws.cell(row=total_row, column=1, value="Total").font = bold_arial
ws.cell(row=total_row, column=2, value=f"=SUM(B6:B{total_row-1})").font = bold_arial
ws.cell(row=total_row, column=3, value=f"=SUM(C6:C{total_row-1})").font = bold_arial
ws.cell(row=total_row, column=4, value=f"=SUM(D6:D{total_row-1})").font = bold_arial

# Formatting
# Currency for price and costs
ws['B1'].number_format = '$#,##0.0000'
for row in range(6, total_row + 1):
    ws.cell(row=row, column=4).number_format = '$#,##0.00'

# 2-decimal precision for gallons
for row in range(6, total_row + 1):
    ws.cell(row=row, column=3).number_format = '0.00'

# Apply Arial to all used cells
for row in ws.iter_rows():
    for cell in row:
        if cell.font.name != 'Arial':
            cell.font = arial_font

# Column widths for readability
ws.column_dimensions['A'].width = 25
ws.column_dimensions['B'].width = 15
ws.column_dimensions['C'].width = 12
ws.column_dimensions['D'].width = 12

wb.save("fuel_costs.xlsx")