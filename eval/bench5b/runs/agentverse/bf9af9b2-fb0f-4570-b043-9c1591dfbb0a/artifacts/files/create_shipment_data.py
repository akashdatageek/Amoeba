import openpyxl
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation

# Create workbook and sheet
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Shipment Data"

# Data
data = [
    ["Lane", "Loads", "Avg Revenue per Load", "Total Revenue"],
    ["Lane A", 12, 1850, "=B2*C2"],
    ["Lane B", 7, 2400, "=B3*C3"],
    ["Lane C", 20, 1420, "=B4*C4"],
    ["Lane D", 4, 3100, "=B5*C5"],
]

# Write data
for row in data:
    ws.append(row)

# Define the Table
# Range A1:D5
tab = Table(displayName="ShipmentTable", ref="A1:D5")

# Add a table style
style = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, 
                       showLastColumn=False, showRowStripes=True, showColumnStripes=False)
tab.tableStyleInfo = style
ws.add_table(tab)

# Totals Row (Row 6)
ws["A6"] = "Total"
ws["B6"] = "=SUM(ShipmentTable[Loads])"
ws["C6"] = "=IFERROR(ROUND(SUM(ShipmentTable[Total Revenue])/SUM(ShipmentTable[Loads]), 2), 0)"
ws["D6"] = "=SUM(ShipmentTable[Total Revenue])"

# Formatting
arial_font = Font(name="Arial")

# Apply Arial font and Currency formatting
for row in ws.iter_rows(min_row=1, max_row=6, min_col=1, max_col=4):
    for cell in row:
        cell.font = arial_font
        if cell.column in [3, 4]:
            cell.number_format = '$#,##0.00'

# Data Validation for Loads (B2:B5)
# Positive integers (whole numbers > 0)
dv = DataValidation(type="whole", operator="greaterThan", formula1="0")
dv.errorTitle = "Invalid Input"
dv.error = "Please enter a positive integer for the number of loads."
ws.add_data_validation(dv)
dv.add("B2:B5")

# Save the file
wb.save("shipment_data.xlsx")
