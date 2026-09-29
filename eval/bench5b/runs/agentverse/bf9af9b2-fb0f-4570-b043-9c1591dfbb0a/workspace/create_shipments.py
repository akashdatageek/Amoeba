import openpyxl
from openpyxl.styles import Font, NumberFormat
from openpyxl.worksheet.datavalidation import DataValidation

# Create workbook and worksheet
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Shipment Data"

# Define data
headers = ["Lane", "Loads", "Avg Revenue per Load", "Total Revenue"]
data = [
    ["Lane A", 12, 1850],
    ["Lane B", 7, 2400],
    ["Lane C", 20, 1420],
    ["Lane D", 4, 3100],
]

# Write headers
ws.append(headers)

# Write data and add row-level formulas for Total Revenue
for i, row_data in enumerate(data, start=2):
    ws.append(row_data)
    # Total Revenue = Loads * Avg Revenue (Column B * Column C)
    ws[f"D{i}"] = f"=B{i}*C{i}"

# Define totals row index
totals_row = len(data) + 2

# Totals Row Labels and Formulas
ws[f"A{totals_row}"] = "Totals"
ws[f"B{totals_row}"] = f"=SUM(B2:B{totals_row-1})"  # Total Loads
ws[f"D{totals_row}"] = f"=SUM(D2:D{totals_row-1})"  # Total Revenue

# Weighted Average Revenue per Load (Total Revenue / Total Loads)
# We use ROUND to ensure 2 decimal places as requested by QA
ws[f"C{totals_row}"] = f"=ROUND(D{totals_row}/B{totals_row}, 2)"

# --- Formatting ---

# Professional Font (Arial)
arial_font = Font(name='Arial')
for row in ws.iter_rows():
    for cell in row:
        cell.font = arial_font

# Currency Formatting for Revenue columns (C and D)
currency_format = '$#,##0.00'
for row in ws.iter_rows(min_row=2, min_col=3, max_col=4):
    for cell in row:
        cell.number_format = currency_format

# Data Validation for 'Loads' column (B2:B5) - Positive Integers
# type='whole', operator='greaterThan', formula='0'
dv = DataValidation(type="whole", operator="greaterThan", formula="0")
ws.add_data_validation(dv)
dv.add("B2:B5")

# Save the file
wb.save("shipment_data.xlsx")