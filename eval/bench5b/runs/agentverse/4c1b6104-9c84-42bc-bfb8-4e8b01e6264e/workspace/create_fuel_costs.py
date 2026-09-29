import openpyxl
from openpyxl.styles import Font

# Data
legs = [
    ("Chicago–Indianapolis", 185),
    ("Indianapolis–Columbus", 175),
    ("Columbus–Pittsburgh", 185)
]
efficiency = 6.5
price = 6.680
effective_date = "09/21/2026"
source_url = "https://www.eia.gov/petroleum/gasdiesel/"

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Fuel Costs"

# Professional Font
arial_font = Font(name='Arial')

# 1. Dynamic Input Section
ws['F1'] = "Fuel Efficiency (MPG)"
ws['G1'] = efficiency
ws['F2'] = "Avg PADD 2 Diesel Price ($/gal)"
ws['G2'] = price
ws['F3'] = "Price Effective Date"
ws['G3'] = effective_date
ws['F4'] = "Source"
ws['G4'] = source_url

# 2. Table Headers
headers = ["Leg", "Distance (miles)", "Gallons", "Cost ($)"]
for col_num, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_num)
    cell.value = header
    cell.font = Font(name='Arial', bold=True)

# 3. Table Data
for row_num, (leg, dist) in enumerate(legs, 2):
    ws.cell(row=row_num, column=1, value=leg)
    ws.cell(row=row_num, column=2, value=dist)
    
    # Gallons: Distance / Efficiency
    gal_cell = ws.cell(row=row_num, column=3)
    gal_cell.value = f"=B{row_num}/$G$1"
    gal_cell.number_format = '0.00'
    
    # Cost: Gallons * Price
    cost_cell = ws.cell(row=row_num, column=4)
    cost_cell.value = f"=C{row_num}*$G$2"
    cost_cell.number_format = '$#,##0.00'

# 4. Total Row
total_row = 5
ws.cell(row=total_row, column=1, value="Total").font = Font(name='Arial', bold=True)

# Total Distance
dist_total = ws.cell(row=total_row, column=2)
dist_total.value = f"=SUM(B2:B4)"
dist_total.font = Font(name='Arial', bold=True)

# Total Gallons
gal_total = ws.cell(row=total_row, column=3)
gal_total.value = f"=SUM(C2:C4)"
gal_total.font = Font(name='Arial', bold=True)
gal_total.number_format = '0.00'

# Total Cost
cost_total = ws.cell(row=total_row, column=4)
cost_total.value = f"=SUM(D2:D4)"
cost_total.font = Font(name='Arial', bold=True)
cost_total.number_format = '$#,##0.00'

# Final formatting for input price
ws['G2'].number_format = '$#,##0.000'

# Apply Arial to all used cells
for row in ws.iter_rows():
    for cell in row:
        if not cell.font.name or cell.font.name != 'Arial':
            cell.font = arial_font

wb.save("fuel_costs.xlsx")
