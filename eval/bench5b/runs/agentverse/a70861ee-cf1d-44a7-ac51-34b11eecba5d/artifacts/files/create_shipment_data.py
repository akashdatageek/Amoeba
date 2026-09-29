import openpyxl
from openpyxl.styles import Font, Protection
from openpyxl.worksheet.datavalidation import DataValidation

# Create workbook and sheet
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Shipment Data"

# Data
data = [
    ["Lane A", 12, 1850],
    ["Lane B", 7, 2400],
    ["Lane C", 20, 1420],
    ["Lane D", 4, 3100],
]

# Headers
headers = ["Lane", "Loads", "Avg Revenue", "Total Revenue"]
for col_num, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_num, value=header)
    cell.font = Font(name="Arial", bold=True)

# Fill data and add formulas for Total Revenue
for row_num, row_data in enumerate(data, 2):
    ws.cell(row=row_num, column=1, value=row_data[0]).font = Font(name="Arial")
    ws.cell(row=row_num, column=2, value=row_data[1]).font = Font(name="Arial")
    ws.cell(row=row_num, column=3, value=row_data[2]).font = Font(name="Arial")
    
    # Total Revenue formula: Loads * Avg Revenue
    formula = f"=B{row_num}*C{row_num}"
    cell = ws.cell(row=row_num, column=4, value=formula)
    cell.font = Font(name="Arial")

# Totals Row
ws.cell(row=6, column=1, value="Totals").font = Font(name="Arial", bold=True)

# Total Loads
ws.cell(row=6, column=2, value="=SUM(B2:B5)").font = Font(name="Arial", bold=True)

# Weighted Average Revenue per Load
# Formula: IFERROR(ROUND(Total Revenue / Total Loads, 2), 0)
ws.cell(row=6, column=3, value="=IFERROR(ROUND(D6/B6, 2), 0)").font = Font(name="Arial", bold=True)

# Total Revenue Sum
ws.cell(row=6, column=4, value="=SUM(D2:D5)").font = Font(name="Arial", bold=True)

# Source Note
ws.cell(row=8, column=1, value="Source: User provided").font = Font(name="Arial")

# Currency Formatting
currency_format = '$#,##0.00'
for row in range(2, 7):
    ws.cell(row=row, column=3).number_format = currency_format
    ws.cell(row=row, column=4).number_format = currency_format

# Data Validation
# Loads: Whole Number >= 0
val_loads = DataValidation(type="whole", operator="greaterThanOrEqual", formula="0", allow_blank=True)
ws.add_data_validation(val_loads)
val_loads.add("B2:B5")

# Avg Revenue: Decimal >= 0
val_rev = DataValidation(type="decimal", operator="greaterThanOrEqual", formula="0", allow_blank=True)
ws.add_data_validation(val_rev)
val_rev.add("C2:C5")

# Protection
# By default, all cells are locked. We unlock the input cells.
for row in range(2, 6):
    ws.cell(row=row, column=2).protection = Protection(locked=False)
    ws.cell(row=row, column=3).protection = Protection(locked=False)

# Protect the sheet to enable the locking
ws.protection.sheet = True
ws.protection.enable()

wb.save("shipment_data.xlsx")