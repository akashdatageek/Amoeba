import openpyxl
from openpyxl.styles import Font

# Create a new workbook and select the active sheet
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Shipment Data"

# Define headers
headers = ["Lane", "Loads", "Avg Revenue", "Total Revenue"]
ws.append(headers)

# Define data
data = [
    ["Lane A", 12, 1850],
    ["Lane B", 7, 2400],
    ["Lane C", 20, 1420],
    ["Lane D", 4, 3100],
]

# Add data and total revenue formulas
for row_idx, row_data in enumerate(data, start=2):
    ws.cell(row=row_idx, column=1, value=row_data[0])
    ws.cell(row=row_idx, column=2, value=row_data[1])
    ws.cell(row=row_idx, column=3, value=row_data[2])
    # Total Revenue = Loads * Avg Revenue
    ws.cell(row=row_idx, column=4, value=f"=B{row_idx}*C{row_idx}")

# Totals row (Row 6)
totals_row = 6
ws.cell(row=totals_row, column=1, value="Total")
ws.cell(row=totals_row, column=2, value=f"=SUM(B2:B{totals_row-1})")
ws.cell(row=totals_row, column=4, value=f"=SUM(D2:D{totals_row-1})")

# Formula for average revenue per load (Weighted Average)
# Avg Revenue = Total Revenue / Total Loads
ws.cell(row=totals_row, column=3, value=f"=D{totals_row}/B{totals_row}")

# Apply professional font (Arial)
arial_font = Font(name="Arial")
for row in ws.iter_rows():
    for cell in row:
        cell.font = arial_font

# Save the workbook
file_path = "shipment_data.xlsx"
wb.save(file_path)
print(f"File saved as {file_path}")
