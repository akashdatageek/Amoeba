from openpyxl import Workbook
from openpyxl.styles import Font

def create_shipment_spreadsheet():
    wb = Workbook()
    ws = wb.active
    ws.title = "Shipment Data"

    # 1. Define Headers
    headers = ["Lane", "Loads", "Avg Revenue", "Total Revenue"]
    ws.append(headers)

    # 2. Raw Data
    data = [
        ("Lane A", 12, 1850),
        ("Lane B", 7, 2400),
        ("Lane C", 20, 1420),
        ("Lane D", 4, 3100),
    ]

    # 3. Populate Data with Dynamic Formulas for Lane Totals
    for i, row_data in enumerate(data, start=2):
        ws.cell(row=i, column=1, value=row_data[0]) # Lane
        ws.cell(row=i, column=2, value=row_data[1]) # Loads
        ws.cell(row=i, column=3, value=row_data[2]) # Avg Revenue
        # Formula: Loads * Avg Revenue
        ws.cell(row=i, column=4, value=f"=B{i}*C{i}")

    # 4. Totals Row
    total_row = 6
    ws.cell(row=total_row, column=1, value="Total")
    ws.cell(row=total_row, column=2, value=f"=SUM(B2:B5)")
    ws.cell(row=total_row, column=4, value=f"=SUM(D2:D5)")

    # 5. Average Revenue per Load Row
    avg_row = 7
    ws.cell(row=avg_row, column=1, value="Avg Revenue/Load")
    # Formula: Total Revenue / Total Loads
    ws.cell(row=avg_row, column=4, value=f"=D{total_row}/B{total_row}")

    # 6. Source Note
    ws.cell(row=9, column=1, value="Source: User provided data")

    # 7. Professional Formatting (Arial Font & Currency)
    arial_font = Font(name='Arial')
    currency_format = '$#,##0.00'

    for row in ws.iter_rows():
        for cell in row:
            cell.font = arial_font
            # Apply currency format to Avg Revenue (Col 3) and Total Revenue (Col 4)
            if cell.column >= 3:
                cell.number_format = currency_format

    # Save the file
    wb.save("shipment_data.xlsx")

if __name__ == "__main__":
    create_shipment_spreadsheet()