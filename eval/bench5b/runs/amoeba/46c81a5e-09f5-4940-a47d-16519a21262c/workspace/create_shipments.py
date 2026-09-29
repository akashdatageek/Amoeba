import openpyxl
from openpyxl.styles import Font

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Shipments"

# Headers
headers = ['Lane', 'Loads', 'Avg Revenue per Load', 'Total Revenue']
ws.append(headers)

# Data
data = [
    ['Lane A', 12, 1850],
    ['Lane B', 7, 2400],
    ['Lane C', 20, 1420],
    ['Lane D', 4, 3100],
]

# Professional Font
arial_font = Font(name='Arial')

# Write data and per-lane total formulas
for i, row_data in enumerate(data, start=2):
    ws.cell(row=i, column=1, value=row_data[0])
    ws.cell(row=i, column=2, value=row_data[1])
    ws.cell(row=i, column=3, value=row_data[2])
    # Total Revenue = Loads * Avg Revenue
    ws.cell(row=i, column=4, value=f'=B{i}*C{i}')

# Totals Row
totals_row = 6
ws.cell(row=totals_row, column=1, value='Totals')
ws.cell(row=totals_row, column=2, value=f'=SUM(B2:B5)')
ws.cell(row=totals_row, column=4, value=f'=SUM(D2:D5)')
# Weighted Average = Total Revenue / Total Loads
ws.cell(row=totals_row, column=3, value=f'=D{totals_row}/B{totals_row}')

# Formatting
currency_format = '$#,##0'
for row in ws.iter_rows(min_row=2, max_row=totals_row, min_col=3, max_col=4):
    for cell in row:
        cell.number_format = currency_format

# Apply Arial font to all used cells
for row in ws.iter_rows(min_row=1, max_row=totals_row, min_col=1, max_col=4):
    for cell in row:
        cell.font = arial_font

wb.save('shipment_data.xlsx')
