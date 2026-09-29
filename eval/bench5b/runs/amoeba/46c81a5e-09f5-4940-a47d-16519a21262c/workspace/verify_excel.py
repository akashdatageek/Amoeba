import openpyxl

file_path = 'shipment_data.xlsx'

# Load with formulas
wb_formula = openpyxl.load_workbook(file_path, data_only=False)
ws_formula = wb_formula.active

# Load with values
wb_value = openpyxl.load_workbook(file_path, data_only=True)
ws_value = wb_value.active

print("--- Formulas and Values ---")
for row in ws_formula.iter_rows():
    row_data = []
    for cell in row:
        val = ws_value[cell.coordinate].value
        form = cell.value if isinstance(cell.value, str) and cell.value.startswith('=') else cell.value
        row_data.append(f"{cell.coordinate}: [Val: {val}, Form: {form}]")
    print(" | ".join(row_data))
