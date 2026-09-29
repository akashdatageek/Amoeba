import openpyxl
from openpyxl.styles import Font

# Create workbook and sheet
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Shipment Data"

# Define headers
headers = ['Lane', 'Loads', 'Avg Revenue per Load', 'Total Revenue']
ws.append(headers)

# Data
data = [
    ['Lane A', 12, 1850],
    ['Lane B', 7, 2400],
    ['Lane C', 20, 1420],
    ['Lane D', 4, 3100],
]

# Professional font
font_arial = Font(name='Arial')

# Write headers with font
for cell in ws[1]:
    cell.font = font_arial

# Write data and add Total Revenue formulas
for i, row_data in enumerate(data, start=2):
    ws.cell(row=i, column=1, value=row_data[0]).font = font_arial
    ws.cell(row=i, column=2, value=row_data[1]).font = font_arial
    ws.cell(row=i, column=3, value=row_data[2]).font = font_arial
    # Total Revenue = Loads * Avg Revenue
    ws.cell(row=i, column=4, value=f'=B{i}*C{i}').font = font_arial

# Totals row (Row 6)
ws.cell(row=6, column=1, value='Totals').font = font_arial
ws.cell(row=6, column=2, value='=SUM(B2:B5)').font = font_arial
# Weighted Average = Total Revenue Sum / Total Loads Sum
ws.cell(row=6, column=3, value='=SUM(D2:D5)/SUM(B2:B5)').font = font_arial
ws.cell(row=6, column=4, value='=SUM(D2:D5)').font = font_arial

# Formatting
currency_format = '$#,##0'
for row in range(2, 7):
    ws.cell(row=row, column=3).number_format = currency_format
    ws.cell(row=row, column=4).number_format = currency_format

# Adjust column widths for readability
ws.column_dimensions['A'].width = 15
ws.column_dimensions['B'].width = 10
ws.column_dimensions['C'].width = 20
ws.column_dimensions['D'].width = 20

wb.save('shipment_data.xlsx')
print("File shipment_data.xlsx created successfully.")