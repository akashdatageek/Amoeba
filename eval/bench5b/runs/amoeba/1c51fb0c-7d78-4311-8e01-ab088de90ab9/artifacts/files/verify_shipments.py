import openpyxl

def verify():
    try:
        # Load workbook to check formulas
        wb = openpyxl.load_workbook('shipment_data.xlsx', data_only=False)
        sheet = wb.active
        
        print("--- Cell Contents and Formulas ---")
        for row in sheet.iter_rows():
            row_data = []
            for cell in row:
                row_data.append(f"{cell.coordinate}: {cell.value}")
            print(" | ".join(row_data))
        
        # Load workbook to check calculated values
        wb_data = openpyxl.load_workbook('shipment_data.xlsx', data_only=True)
        sheet_data = wb_data.active
        
        print("\n--- Calculated Values ---")
        for row in sheet_data.iter_rows():
            row_data = []
            for cell in row:
                row_data.append(f"{cell.coordinate}: {cell.value}")
            print(" | ".join(row_data))
            
    except Exception as e:
        print(f"Error: {e}")

if __name__ == "__main__":
    verify()
