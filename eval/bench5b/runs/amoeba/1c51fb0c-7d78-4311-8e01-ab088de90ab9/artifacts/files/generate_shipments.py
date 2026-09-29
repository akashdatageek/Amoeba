import openpyxl
from openpyxl.styles import Font

# Create a new workbook and select the active sheet
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Shipment Data"

# 1. Create table headers
headers = ['Lane', 'Loads', 'Avg Revenue per Load', 'Total Revenue']
ws.append(headers)

# Set header font to bold
for cell in ws[1]:
    cell.font = Font(bold=True)

# Data to be inserted
data = [
    ('Lane A', 12, 1850),
    ('Lane B', 7, 2400),
    ('Lane C', 20, 1420),
    ('Lane D', 4, 3100),
]

# 2. Populate data and calculate Total Revenue per lane using formulas
for i, (lane, loads, avg) in enumerate(data, start=2):
    ws.cell(row=i, column=1, value=lane)
    ws.cell(row=i, column=2, value=loads)
    ws.cell(row=i, column=3, value=avg)
    # Formula for Total Revenue: Loads * Avg Revenue
    ws.cell(row=i, column=4, value=f'=B{i}*C{i}')

# 3. Add a 'Totals' row
totals_row = 6
ws.cell(row=totals_row, column=1, value='Totals').font = Font(bold=True)

# 4. Use SUM formulas for totals
# Total Loads
ws.cell(row=totals_row, column=2, value=f'=SUM(B2:B5)').font = Font(bold=True)
# Total Revenue
ws.cell(row=totals_row, column=4, value=f'=SUM(D2:D5)').font = Font(bold=True)

# 5. Create a separate row for Weighted Average Revenue per Load
avg_row = 7
ws.cell(row=avg_row, column=1, value='Weighted Avg Revenue per Load').font = Font(bold=True)
# Formula: Total Revenue Sum / Total Loads Sum
ws.cell(row=avg_row, column=3, value='=SUM(D2:D5)/SUM(B2:B5)').font = Font(bold=True)

# Formatting: Currency for Avg Revenue and Total Revenue columns
currency_format = '$#,##0.00'
for row in range(2, 8):
    # Column C (Avg Revenue)
    ws.cell(row=row, column=3).number_format = currency_format
    # Column D (Total Revenue)
    ws.cell(row=row, column=4).number_format = currency_format

# Adjust column widths for readability
ws.column_dimensions['A'].width = 25
ws.column_dimensions['B'].width = 10
ws.column_dimensions['C'].width = 20
ws.column_dimensions['D'].width = 20

# 6. Save the file
wb.save('shipment_data.xlsx')
print("File shipment_data.xlsx created successfully.")