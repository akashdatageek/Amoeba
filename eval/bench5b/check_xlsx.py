"""bench5b — open every .xlsx a run made (a copy, recalculated with LibreOffice headless, D71) and print each sheet's
cells: formula and computed value side by side, so the observer can check the file tasks (bench5-xlsx, bench5-diesel).
Read-only on the run folders.

    python eval/bench5b/check_xlsx.py [task] [--rep k]
"""
import json
import shutil
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
from amoeba.localtools.office import recalc  # noqa: E402


def show(path: Path) -> None:
    with tempfile.TemporaryDirectory() as tmp:
        copy = Path(tmp) / path.name
        shutil.copy(path, copy)
        rc = recalc(copy)
        f, v = load_workbook(copy), load_workbook(copy, data_only=True)
        print(f"  file {path.name} (recalc: {rc['detail']})")
        for ws in f.worksheets:
            print(f"    sheet {ws.title!r}")
            for row in ws.iter_rows():
                cells = []
                for c in row:
                    if c.value is None:
                        continue
                    val = v[ws.title][c.coordinate].value
                    cells.append(f"{c.coordinate}={c.value!r}" + (f"→{val!r}" if str(c.value).startswith("=") else ""))
                if cells:
                    print("      " + "  ".join(cells))


def main() -> None:
    task = next((a for a in sys.argv[1:] if not a.startswith("--") and not a.isdigit()), None)
    rep = int(sys.argv[sys.argv.index("--rep") + 1]) if "--rep" in sys.argv else None
    for line in (ROOT / "eval/bench5b/ledger.jsonl").read_text().splitlines():
        r = json.loads(line)
        if (task and r["task"] != task) or (rep and r["rep"] != rep) or not r.get("run"):
            continue
        d = ROOT / r["run"]
        if not d.exists():
            d = ROOT / "eval/bench5b/runs" / r["arch"] / Path(r["run"]).name
        files = [p for p in d.rglob("*.xlsx") if "skills" not in p.parts and "artifacts" not in p.parts]
        print(f"== {r['task']} {r['arch']} r{r['rep']}: {len(files)} xlsx")
        for p in files:
            show(p)


if __name__ == "__main__":
    main()
