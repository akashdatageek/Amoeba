import openpyxl
from openpyxl.styles import Font

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Cost Plan"

# Professional Font
arial_font = Font(name='Arial')

# 1. Load Data
loads_data = [
    ["Load ID", "Weight (lb)", "Miles"],
    ["L1", 18000, 220],
    ["L2", 12500, 340],
    ["L3", 22000, 150],
    ["L4", 9000, 410],
    ["L5", 15500, 95],
    ["L6", 7000, 280],
    ["L7", 20500, 190],
    ["L8", 11000, 365],
    ["L9", 14000, 120],
    ["L10", 8500, 255],
]
for r_idx, row in enumerate(loads_data, 1):
    for c_idx, value in enumerate(row, 1):
        cell = ws.cell(row=r_idx, column=c_idx, value=value)
        cell.font = arial_font

# 2. Truck Data
trucks_data = [
    ["Truck ID", "Capacity (lb)", "Limit (mi)", "Fixed Cost ($)", "Rate ($/mi)"],
    ["A", 45000, 1100, 350, 2.10],
    ["B", 45000, 1100, 350, 2.10],
    ["C", 30000, 1100, 500, 1.85],
    ["D", 26000, 900, 275, 2.45],
]
for r_idx, row in enumerate(trucks_data, 1):
    for c_idx, value in enumerate(row, 5): # Start at column E
        cell = ws.cell(row=r_idx, column=c_idx, value=value)
        cell.font = arial_font

# 3. Assignment & Cost Table
assignment_headers = ["Truck", "Assigned Loads", "Total Weight (lb)", "Total Miles", "Cost ($)"]
for c_idx, header in enumerate(assignment_headers, 11): # Start at column K
    cell = ws.cell(row=1, column=c_idx, value=header)
    cell.font = Font(name='Arial', bold=True)

# Assignments based on Step 1 [S1]
# Truck A: L1, L2, L9 -> Weights: B2, B3, B10; Miles: C2, C3, C10
# Truck B: L5, L7, L10 -> Weights: B6, B8, B11; Miles: C6, C8, C11
# Truck C: L4, L6, L8 -> Weights: B5, B7, B9; Miles: C5, C7, C9
# Truck D: L3 -> Weights: B4; Miles: C4

assignments = [
    ["A", "L1, L2, L9", "=B2+B3+B10", "=C2+C3+C10", "=H2+(N2*I2)"],
    ["B", "L5, L7, L10", "=B6+B8+B11", "=C6+C8+C11", "=H3+(N3*I3)"],
    ["C", "L4, L6, L8", "=B5+B7+B9", "=C5+C7+C9", "=H4+(N4*I4)"],
    ["D", "L3", "=B4", "=C4", "=H5+(N5*I5)"],
]

for r_idx, row in enumerate(assignments, 2):
    for c_idx, value in enumerate(row, 11):
        cell = ws.cell(row=r_idx, column=c_idx, value=value)
        cell.font = arial_font

# Total Cost
ws.cell(row=6, column=14, value="Total Cost:").font = arial_font
ws.cell(row=6, column=15, value="=SUM(O2:O5)").font = arial_font

wb.save("cost_plan_xlsx.xlsx")
