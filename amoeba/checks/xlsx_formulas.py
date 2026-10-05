"""D113 — spreadsheet checks: totals and derived cells must be formulas, not typed values.

For each .xlsx/.xlsm file a step made, plain code opens the workbook with its formulas (openpyxl) and lists the typed
numbers that should have been formulas: a number in a total row or column (its row label or column header says total,
sum, subtotal, grand total, average, mean or net), a number equal to the sum of the two or more numbers just left of
it in its row or just above it in its column, and a number equal to the product of two numbers to its left in its
row. A number the task states is an input, never flagged. Any such cell fails the check (the step gets the retry turn)
with the cell names, so the workbook follows its inputs when they change. (Probe (hard) H1: the per-truck weights and
miles were typed in.)
"""
from __future__ import annotations

import re
from pathlib import Path

TOTAL = re.compile(r"\b(?:total|totals|sum|subtotal|grand total|average|mean|net)\b", re.I)
MAX_CELLS = 20


def _num(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _formula(cell) -> bool:
    return cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("="))


def _text(v) -> bool:
    """A label: text that is not a formula."""
    return isinstance(v, str) and not v.startswith("=")


def _same(a: float, b: float) -> bool:
    return abs(a - b) <= max(0.01, 1e-6 * abs(b))


def _given(task: str) -> set[float]:
    out = set()
    for m in re.finditer(r"(?<![\w.])\d(?:[\d,]*\d)?(?:\.\d+)?", task or ""):
        try:
            out.add(round(float(m.group(0).replace(",", "")), 6))
        except ValueError:
            pass
    return out


# box: niche
def typed_derived_cells(path: str | Path, task: str = "") -> list[dict]:
    """Typed numbers where a formula belongs: {sheet, cell, value, why}."""
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    wb = load_workbook(path, data_only=False)
    given = _given(task)
    out = []
    for ws in wb.worksheets:
        grid = {(c.row, c.column): c for row in ws.iter_rows() for c in row if c.value is not None}

        def typed(r, c):
            x = grid.get((r, c))
            return x is not None and _num(x.value) and not _formula(x)

        def numeric(r, c):
            x = grid.get((r, c))
            return x is not None and (_num(x.value) or _formula(x))

        def label(r, c):
            row = next((grid[(r, k)].value for k in range(1, c)
                        if (r, k) in grid and _text(grid[(r, k)].value)), "")
            col = next((grid[(k, c)].value for k in range(r - 1, 0, -1)
                        if (k, c) in grid and _text(grid[(k, c)].value)), "")
            return f"{row} {col}"

        for (r, c), cell in sorted(grid.items()):
            if not typed(r, c) or round(float(cell.value), 6) in given:
                continue
            v, why = float(cell.value), None
            name = f"{ws.title}!{get_column_letter(c)}{r}"
            left = []
            k = c - 1
            while k >= 1 and typed(r, k):
                left.append(float(grid[(r, k)].value))
                k -= 1
            up = []
            k = r - 1
            while k >= 1 and typed(k, c):
                up.append(float(grid[(k, c)].value))
                k -= 1
            if TOTAL.search(label(r, c)):
                why = "a total typed in"
            elif len(left) >= 2 and _same(sum(left), v):
                why = f"the sum of the {len(left)} cells to its left"
            elif len(up) >= 2 and _same(sum(up), v):
                why = f"the sum of the {len(up)} cells above it"
            else:
                near = left[:4]
                for i in range(len(near)):
                    for j in range(i + 1, len(near)):
                        a, b = near[i], near[j]
                        if a not in (0, 1) and b not in (0, 1) and _same(a * b, v) and not _same(v, a):
                            why = "the product of two cells to its left"
            if why:
                out.append({"sheet": ws.title, "cell": name, "value": cell.value, "why": why})
            if len(out) >= MAX_CELLS:
                return out
    return out


def check(output: str, evidence: dict) -> tuple[bool, str]:
    """Pass when no workbook the step made has a typed total or derived cell."""
    found = []
    for f in evidence.get("xlsx_files") or []:
        try:
            found += typed_derived_cells(f, evidence.get("task", ""))
        except Exception as e:                        # an unreadable workbook is reported, not a crash
            return False, f"the workbook {Path(f).name} could not be opened to check its formulas ({type(e).__name__})"
    if not found:
        return True, "every total and derived cell is a formula"
    cells = "; ".join(f"{x['cell']} = {x['value']} ({x['why']})" for x in found[:10])
    return False, (f"these spreadsheet cells are typed numbers where a formula belongs: {cells}. Write each as a "
                   f"formula (=SUM(…), =B5*C5, =SUMIF(…)) that references the input cells, so the workbook follows "
                   f"its inputs, and save the file again.")
