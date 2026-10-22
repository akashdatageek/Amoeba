"""D142 — the plan–execution contract: typed step checks from a fixed vocabulary that plain code evaluates.

Box 2 writes, per step, what "done" means as `checks` (each with the requirement ids it `covers`); `done_when` stays
one sentence for the reviewers. Box 3 runs a step's checks on its output text and its workspace files after the
step's Final Output (a failure is a failed check like any other: the check-retry turns, then partial or incomplete,
and Box 4's `checks` cause). The final requirement check counts a requirement as met only when every check covering
it passed. A plan without checks behaves as before: the checks are derived from the output line's format markers.
"""
from __future__ import annotations

import csv
import io
import math
import re
from pathlib import Path

FILE_KINDS = ("docx", "xlsx", "csv", "png", "pdf", "pptx", "md", "json")
# name → {param: (type, required)}; types: int, str, kind, list (of strings), image kind
VOCAB: dict[str, dict[str, tuple[str, bool]]] = {
    "file_exists": {"kind": ("kind", True), "name": ("str", False)},
    "docx_pages_max": {"n": ("int", True)},
    "docx_words_between": {"min": ("int", True), "max": ("int", True)},
    "docx_has_table": {"min_rows": ("int", False)},
    "xlsx_rows_min": {"n": ("int", True), "sheet": ("str", False)},
    "xlsx_columns": {"names": ("list", True)},
    "xlsx_has_formulas": {},
    "csv_rows_min": {"n": ("int", True)},
    "csv_columns": {"names": ("list", True)},
    "image_exists": {"kind": ("image", True), "min_px": ("int", False)},
    "output_has_table": {"min_rows": ("int", False)},
    "output_has_list": {"min_items": ("int", False)},
    "output_mentions": {"terms": ("list", True)},
    "cites_every_figure": {},
    "figures_dated": {},
    "verdict_present": {},
    "not_sent_fallback": {},
}
FILE_CHECKS = {"file_exists", "docx_pages_max", "docx_words_between", "docx_has_table", "xlsx_rows_min",
               "xlsx_columns", "xlsx_has_formulas", "csv_rows_min", "csv_columns", "image_exists", "not_sent_fallback"}
MAX_CHECKS = 6                      # checks_reachable: more than this on one step is unlikely to be met in its turns
WORDS_PER_PAGE = 500


# box: checks
def check_problem(c) -> str | None:
    """Why a check is not in the vocabulary or its parameters are malformed (None when it is valid)."""
    if not isinstance(c, dict):
        return f"not an object: {str(c)[:60]}"
    name, params = c.get("name"), c.get("params") or {}
    if name not in VOCAB:
        return f"unknown check {name!r} (allowed: {', '.join(VOCAB)})"
    if not isinstance(params, dict):
        return f"{name}: params must be an object"
    spec = VOCAB[name]
    extra = sorted(set(params) - set(spec))
    if extra:
        return f"{name}: unknown parameter(s) {', '.join(extra)}"
    for p, (typ, required) in spec.items():
        v = params.get(p)
        if v is None:
            if required:
                return f"{name}: missing parameter {p}"
            continue
        if typ == "int" and not (isinstance(v, int) and not isinstance(v, bool) and v >= 0):
            return f"{name}: {p} must be a whole number >= 0"
        if typ == "str" and not (isinstance(v, str) and v.strip()):
            return f"{name}: {p} must be a non-empty string"
        if typ == "kind" and v not in FILE_KINDS:
            return f"{name}: kind must be one of {', '.join(FILE_KINDS)}"
        if typ == "image" and v not in ("png", "svg"):
            return f"{name}: kind must be png or svg"
        if typ == "list" and not (isinstance(v, list) and v and all(isinstance(x, str) and x.strip() for x in v)):
            return f"{name}: {p} must be a non-empty list of strings"
    if name == "docx_words_between" and params["min"] > params["max"]:
        return f"{name}: min is above max"
    covers = c.get("covers", [])
    if not isinstance(covers, list) or not all(isinstance(r, str) and re.fullmatch(r"R\d+", r) for r in covers):
        return f"{name}: covers must be a list of requirement ids (R1, R2, ...)"
    return None


# box: checks
def parse_checks_line(raw: str) -> list[dict]:
    """A legacy plan's optional `checks:` line: a JSON list of {name, params, covers}. Unreadable JSON is kept as one
    `unparsed` check, so checks_valid reports it to the planner instead of dropping it."""
    import json
    try:
        v = json.loads(raw)
    except ValueError:
        return [{"name": "unparsed", "params": {}, "covers": [], "raw": raw[:200]}]
    v = v if isinstance(v, list) else [v]
    return [dict(x, params=x.get("params") or {}, covers=x.get("covers") or []) if isinstance(x, dict) else
            {"name": "unparsed", "params": {}, "covers": [], "raw": str(x)[:200]} for x in v]


# box: checks
def derived_checks(output: str, covers: list[str]) -> list[dict]:
    """A step without checks: the ones its output line's format markers imply (what Box 3 checks today), so a
    legacy plan is judged as before."""
    from amoeba.interp.plan_runner import output_markers
    marks = output_markers(output or "")
    out = []
    if "table" in marks:
        out.append({"name": "output_has_table", "params": {"min_rows": 1}, "covers": list(covers), "derived": True})
    if "list" in marks:
        out.append({"name": "output_has_list", "params": {"min_items": 2}, "covers": list(covers), "derived": True})
    return out


# box: plan_step
def checks_text(checks: list[dict]) -> str:
    """The step's checks in one line for its helpers: 'docx_pages_max (n=2); output_mentions (terms=[...])'."""
    def one(c):
        p = ", ".join(f"{k}={v}" for k, v in (c.get("params") or {}).items())
        return f"{c.get('name')}{f' ({p})' if p else ''}"
    return "; ".join(one(c) for c in checks)


def effective_checks(step) -> list[dict]:
    return list(step.checks) if getattr(step, "checks", None) else derived_checks(step.output, step.covers)


# ---- evaluation ---------------------------------------------------------------------------------------------------
def _files(files: list[str], kind: str, name: str | None = None) -> list[Path]:
    ps = [Path(f) for f in files if Path(f).suffix.lower().lstrip(".") == kind]
    if name:
        ps = [p for p in ps if p.name.lower() == name.lower() or p.stem.lower() == Path(name).stem.lower()]
    return ps


def _docx_stats(p: Path) -> dict:
    import docx
    d = docx.Document(str(p))
    words = sum(len(x.text.split()) for x in d.paragraphs)
    words += sum(len(c.text.split()) for t in d.tables for r in t.rows for c in r.cells)
    breaks = len(d.element.body.xpath('.//w:br[@w:type="page"]'))
    return {"words": words, "pages": breaks + 1 if breaks else max(1, math.ceil(words / WORDS_PER_PAGE)),
            "estimated": not breaks, "table_rows": max((len(t.rows) for t in d.tables), default=0)}


def _xlsx(p: Path, sheet: str | None = None):
    import openpyxl
    wb = openpyxl.load_workbook(str(p))
    return wb[sheet] if sheet and sheet in wb.sheetnames else wb.worksheets[0], wb


def _rows(ws) -> list[list]:
    return [list(r) for r in ws.iter_rows(values_only=True) if any(v not in (None, "") for v in r)]


def _csv_rows(p: Path) -> list[list[str]]:
    text = p.read_text(encoding="utf-8", errors="replace")
    return [r for r in csv.reader(io.StringIO(text)) if any(x.strip() for x in r)]


def _header_has(header: list, names: list[str]) -> list[str]:
    have = {str(h).strip().lower() for h in header if h is not None}
    return [n for n in names if n.strip().lower() not in have]


def table_rows(text: str) -> int:
    """Data rows of the largest markdown table in the text (header and |---| row not counted)."""
    best, run = 0, []
    for line in (text or "").splitlines() + [""]:
        if re.match(r"^\s*\|.*\|\s*$", line):
            run.append(line)
            continue
        if run:
            seps = [i for i, x in enumerate(run) if re.match(r"^\s*\|?\s*:?-{3,}", x)]
            best = max(best, len(run) - (seps[0] + 1) if seps else 0)
        run = []
    return best


def list_items(text: str) -> int:
    return len(re.findall(r"^\s*(?:[-*]|\d+[.)])\s+\S", text or "", re.M))


# box: step_check
def run_check(c: dict, ctx: dict) -> dict:
    """One check against a step's output (ctx["text"]) and files (ctx["files"]: absolute paths). ctx may also carry
    "untagged" (figures without a source tag, D33), "undated" (web figures without a date, D110) and "verdict"
    (the verify step's parsed verdict). Returns {name, pass, detail, covers, params}."""
    name, p = c.get("name"), c.get("params") or {}
    out = {"name": f"contract:{name}", "check": name, "params": p, "covers": list(c.get("covers") or []),
           "source": "contract"}
    bad = check_problem(c)
    if bad:
        return {**out, "pass": False, "detail": f"invalid check: {bad}"}
    text, files = ctx.get("text") or "", ctx.get("files") or []
    try:
        ok, detail = _evaluate(name, p, text, files, ctx)
    except Exception as e:                       # an unreadable file fails the check, it never crashes the step
        ok, detail = False, f"could not read the file: {type(e).__name__}: {str(e)[:120]}"
    return {**out, "pass": bool(ok), "detail": detail}


def _evaluate(name: str, p: dict, text: str, files: list[str], ctx: dict) -> tuple[bool, str]:
    if name == "file_exists":
        got = _files(files, p["kind"], p.get("name"))
        want = p.get("name") or f"a .{p['kind']} file"
        return bool(got), f"found {got[0].name}" if got else f"{want} is not in the step's files"
    if name.startswith("docx_"):
        docs = _files(files, "docx")
        if not docs:
            return False, "no .docx file in the step's files"
        s = _docx_stats(docs[0])
        est = " (estimated at 500 words per page)" if s["estimated"] else ""
        if name == "docx_pages_max":
            return s["pages"] <= p["n"], f"{docs[0].name}: {s['pages']} page(s){est}, at most {p['n']} allowed"
        if name == "docx_words_between":
            return p["min"] <= s["words"] <= p["max"], f"{docs[0].name}: {s['words']} words, {p['min']}–{p['max']} wanted"
        k = p.get("min_rows") or 1
        return s["table_rows"] >= k, f"{docs[0].name}: largest table has {s['table_rows']} row(s), {k} wanted"
    if name.startswith("xlsx_"):
        books = _files(files, "xlsx")
        if not books:
            return False, "no .xlsx file in the step's files"
        ws, wb = _xlsx(books[0], p.get("sheet"))
        if name == "xlsx_has_formulas":
            n = sum(isinstance(v, str) and v.startswith("=") for s in wb.worksheets
                    for r in s.iter_rows(values_only=True) for v in r)
            return n > 0, f"{books[0].name}: {n} formula cell(s)"
        rows = _rows(ws)
        if name == "xlsx_rows_min":
            data = max(len(rows) - 1, 0)
            return data >= p["n"], f"{books[0].name} [{ws.title}]: {data} data row(s), {p['n']} wanted"
        miss = _header_has(rows[0] if rows else [], p["names"])
        return not miss, f"{books[0].name}: header lacks {', '.join(miss)}" if miss else f"{books[0].name}: header ok"
    if name.startswith("csv_"):
        cs = _files(files, "csv")
        if not cs:
            return False, "no .csv file in the step's files"
        rows = _csv_rows(cs[0])
        if name == "csv_rows_min":
            data = max(len(rows) - 1, 0)
            return data >= p["n"], f"{cs[0].name}: {data} data row(s), {p['n']} wanted"
        miss = _header_has(rows[0] if rows else [], p["names"])
        return not miss, f"{cs[0].name}: header lacks {', '.join(miss)}" if miss else f"{cs[0].name}: header ok"
    if name == "image_exists":
        imgs = _files(files, p["kind"])
        if not imgs:
            return False, f"no .{p['kind']} file in the step's files"
        if p.get("min_px") and p["kind"] == "png":
            from PIL import Image
            with Image.open(imgs[0]) as im:
                w, h = im.size
            return min(w, h) >= p["min_px"], f"{imgs[0].name}: {w}x{h} px, at least {p['min_px']} wanted"
        return True, f"found {imgs[0].name}"
    if name == "output_has_table":
        k, got = p.get("min_rows") or 1, table_rows(text)
        return got >= k, f"the output's largest table has {got} data row(s), {k} wanted"
    if name == "output_has_list":
        k, got = p.get("min_items") or 2, list_items(text)
        return got >= k, f"the output has {got} list item(s), {k} wanted"
    if name == "output_mentions":
        low = text.lower()
        miss = [t for t in p["terms"] if t.lower() not in low]
        return not miss, f"the output does not mention: {', '.join(miss)}" if miss else "all terms mentioned"
    if name == "cites_every_figure":
        n = int(ctx.get("untagged") or 0)
        return n == 0, f"{n} figure(s) without a source tag" if n else "every figure has a source tag or is derived"
    if name == "figures_dated":
        n = len(ctx.get("undated") or [])
        return n == 0, f"{n} web figure(s) without a date" if n else "every web figure is dated"
    if name == "verdict_present":
        v = ctx.get("verdict")
        return v is not None, f"verdict {v}" if v else "no Verdict: PASS or Verdict: FAIL line"
    if name == "not_sent_fallback":
        marked = [f for f in files if re.search(r"not[_ -]?sent", Path(f).name, re.I) or _says_not_sent(f)]
        limits = bool(re.search(r"(?im)^\W*limitations\b", text)) and bool(re.search(r"(?i)not sent", text))
        why = ([] if marked else ["no file marked NOT SENT"]) + ([] if limits else ["no Limitations line saying NOT SENT"])
        return not why, "; ".join(why) or f"{Path(marked[0]).name} marked NOT SENT and named under Limitations"
    return False, f"no evaluator for {name}"


def _says_not_sent(f: str) -> bool:
    try:
        return Path(f).suffix.lower() in (".md", ".txt", ".eml", ".html") and \
            "NOT SENT" in Path(f).read_text(encoding="utf-8", errors="replace")[:4000]
    except OSError:
        return False


# box: step_check
def run_checks(checks: list[dict], ctx: dict) -> list[dict]:
    return [run_check(c, ctx) for c in checks]


# ---- Box 2: the plan's checks (quality.py's HARD checks) -----------------------------------------------------------
FILE_TOOL_HINT = re.compile(r"local[:_](?:write|bash|edit)|^made:|\b(write|file|docx|xlsx|excel|spreadsheet|python|"
                            r"code|bash|chart|plot|csv|pdf|pptx|document)\b", re.I)


# box: checks
def plan_contract_quality(requirements: list[str], steps: list, roles_by_name: dict) -> dict:
    """requirements_checked, checks_valid and checks_reachable for a draft whose steps carry checks (n/a otherwise).
    roles_by_name: name → its tools (installed and requested)."""
    if not any(getattr(s, "checks", None) for s in steps):
        na = {"ok": None, "detail": "no step has checks (legacy plan: checks derived from the format markers)"}
        return {"requirements_checked": na, "checks_valid": dict(na), "checks_reachable": dict(na)}
    eff = {s.index + 1: effective_checks(s) for s in steps}
    covered = {r for cs in eff.values() for c in cs for r in c.get("covers") or []}
    unchecked = [r for r in requirements if r not in covered]
    invalid = [{"step": n, "problem": pb} for n, cs in eff.items() for c in cs if (pb := check_problem(c))]
    many = [{"step": n, "checks": len(cs)} for n, cs in eff.items() if len(cs) > MAX_CHECKS]
    no_writer = []
    for s in steps:
        names = [c.get("name") for c in eff[s.index + 1] if c.get("name") in FILE_CHECKS]
        tools = [t for a in s.agent_names for t in roles_by_name.get(a, [])]
        if names and not any(FILE_TOOL_HINT.search(t) for t in tools):
            no_writer.append({"step": s.index + 1, "checks": names})
    return {"requirements_checked": {"ok": not unchecked, "unchecked": unchecked} if requirements else
            {"ok": None, "detail": "no requirement ids"},
            "checks_valid": {"ok": not invalid, "invalid": invalid},
            "checks_reachable": {"ok": not many and not no_writer, "too_many": many, "no_file_tool": no_writer}}


# ---- the final answer --------------------------------------------------------------------------------------------
LIMITATIONS = re.compile(r"(?im)^\s*(?:#{1,6}\s*|\*\*\s*)?limitations\b")


# box: runresult
def split_limitations(answer: str) -> tuple[str, str]:
    """The answer body above its Limitations section, and that section (empty when there is none)."""
    m = LIMITATIONS.search(answer or "")
    return ((answer or "")[:m.start()], (answer or "")[m.start():]) if m else (answer or "", "")
